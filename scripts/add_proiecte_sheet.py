"""Adauga foaia 'Proiecte' in continut-site.xlsx (fara sa atinga restul
foilor deja existente / editate de utilizator). Se ruleaza o singura data.
"""
import json, os
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = f"{ROOT}/continut-site.xlsx"
SRC = f"{ROOT}/portofoliu poze solar electric panel/deploy/data.js"

text = open(SRC, encoding="utf-8").read()
text = text[text.index("{"):].rstrip().rstrip(";")
data = json.loads(text)

wb = openpyxl.load_workbook(XLSX)
if "Proiecte" in wb.sheetnames:
    del wb["Proiecte"]
ws = wb.create_sheet("Proiecte")
ws.append(["id", "nume", "album", "locatie", "note", "data_de_la", "data_pana", "poze"])
for p in data["projects"]:
    ws.append([p["id"], p["name"], p["album"], p.get("location") or "", p.get("notes") or "", p["date_from"], p["date_to"], p["count"]])

for i in range(1, ws.max_column + 1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 28

wb.save(XLSX)
print(f"Adaugat {len(data['projects'])} proiecte in foaia 'Proiecte' din {XLSX}")
