"""Adauga foaia 'Poze' in continut-site.xlsx (nota + ascundere per poza),
fara sa atinga celelalte foi. Se ruleaza o singura data la migrare.
"""
import json, os
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = f"{ROOT}/continut-site.xlsx"
PHOTOS_SRC = f"{ROOT}/scripts/portfolio_photos_source.json"

photo_src = json.load(open(PHOTOS_SRC, encoding="utf-8"))

wb = openpyxl.load_workbook(XLSX)
if "Poze" in wb.sheetnames:
    del wb["Poze"]
ws = wb.create_sheet("Poze")
ws.append(["id", "project_id", "nota", "ascunsa"])
for p in photo_src["photos"]:
    ws.append([p["id"], p["project_id"], "", "nu"])

for i in range(1, ws.max_column + 1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 22

wb.save(XLSX)
print(f"Adaugat {len(photo_src['photos'])} poze in foaia 'Poze'")
