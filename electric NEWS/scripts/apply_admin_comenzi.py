"""Aplică automat comenzile create din admin/gestionare/ (add/edit/delete articole, întrebări, surse, categorii).

Operații mecanice, deterministe — nu necesită judecată editorială, deci rulează direct,
fără să treacă prin Claude de fiecare dată. Rulare: py scripts\\apply_admin_comenzi.py
"""
import openpyxl, os, json, shutil, datetime, glob
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = f"{ROOT}/electric-news-calendar-editorial.xlsx"
CATEGORII_PATH = f"{ROOT}/categorii.json"
ARTICLE_TEXT_DIR = f"{ROOT}/articole-text"
IMG_ART_DIR = f"{ROOT}/site-local/assets/images/articole"
IMG_MANIFEST_PATH = f"{IMG_ART_DIR}/.image-manifest.json"
COMENZI_DIR = f"{ROOT}/cereri-publicare/comenzi"
COMENZI_DONE_DIR = f"{COMENZI_DIR}/executate"

os.makedirs(COMENZI_DIR, exist_ok=True)
os.makedirs(COMENZI_DONE_DIR, exist_ok=True)

files = sorted(glob.glob(f"{COMENZI_DIR}/*.json"))
if not files:
    print("Nicio comandă nouă în cereri-publicare/comenzi/.")
    raise SystemExit(0)

comenzi = []
for fp in files:
    with open(fp, encoding="utf-8") as f:
        comenzi.append((fp, json.load(f)))

wb = openpyxl.load_workbook(XLSX)
ws = wb["Calendar editorial"]
aqa_ws = wb["Q&A Articole"]
categorii = json.load(open(CATEGORII_PATH, encoding="utf-8"))
img_manifest = json.load(open(IMG_MANIFEST_PATH, encoding="utf-8")) if os.path.exists(IMG_MANIFEST_PATH) else {}


def slugify(s):
    repl = {"ă":"a","â":"a","î":"i","ș":"s","ş":"s","ț":"t","ţ":"t","Ă":"a","Â":"a","Î":"i","Ș":"s","Ț":"t"}
    out = "".join(repl.get(ch, ch) for ch in s)
    out = out.lower()
    keep = []
    for ch in out:
        if ch.isalnum():
            keep.append(ch)
        elif ch in " -_/":
            keep.append(" ")
    out = "".join(keep)
    out = "-".join(out.split())
    return out[:80].strip("-")


def slug_row_map():
    mapping = {}
    seen = {}
    for r in range(2, ws.max_row + 1):
        titlu = ws.cell(row=r, column=5).value
        if not titlu:
            continue
        base = slugify(titlu)
        slug = base
        n = 2
        while slug in seen:
            slug = f"{base}-{n}"
            n += 1
        seen[slug] = True
        mapping[slug] = r
    return mapping


def find_template_row(categorie, tip):
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=3).value == categorie and ws.cell(row=r, column=4).value == tip:
            return r
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=4).value == tip:
            return r
    return 2


def copy_row_style(src_row, dst_row):
    for col in range(1, 11):
        src_cell = ws.cell(row=src_row, column=col)
        dst_cell = ws.cell(row=dst_row, column=col)
        dst_cell.font = Font(**{k: getattr(src_cell.font, k) for k in ["name", "size", "bold", "italic", "color"]})
        if src_cell.fill and src_cell.fill.fgColor and src_cell.fill.fgColor.rgb not in (None, "00000000"):
            dst_cell.fill = PatternFill(fill_type="solid", fgColor=src_cell.fill.fgColor.rgb)
        dst_cell.alignment = Alignment(**{k: getattr(src_cell.alignment, k) for k in ["horizontal", "vertical", "wrapText"]})
    dst_cell_date = ws.cell(row=dst_row, column=2)
    dst_cell_date.number_format = ws.cell(row=src_row, column=2).number_format


def sync_status(slug, row):
    txt_path = f"{ARTICLE_TEXT_DIR}/{slug}.txt"
    ws.cell(row=row, column=8, value="Publicat" if os.path.exists(txt_path) else "În așteptare")


def write_continut(slug, continut):
    if continut and continut.strip():
        with open(f"{ARTICLE_TEXT_DIR}/{slug}.txt", "w", encoding="utf-8") as f:
            f.write(continut.strip() + "\n")


def rename_article_assets(old_slug, new_slug):
    if old_slug == new_slug:
        return
    old_txt = f"{ARTICLE_TEXT_DIR}/{old_slug}.txt"
    if os.path.exists(old_txt):
        shutil.move(old_txt, f"{ARTICLE_TEXT_DIR}/{new_slug}.txt")
    old_img = f"{IMG_ART_DIR}/{old_slug}.jpg"
    if os.path.exists(old_img):
        shutil.move(old_img, f"{IMG_ART_DIR}/{new_slug}.jpg")
    if old_slug in img_manifest:
        img_manifest[new_slug] = img_manifest.pop(old_slug)
    for r in range(2, aqa_ws.max_row + 1):
        if (aqa_ws.cell(row=r, column=1).value or "").strip() == old_slug:
            aqa_ws.cell(row=r, column=1, value=new_slug)


pending_delete_intrebare_rows = set()
pending_delete_articol_slugs = set()
log = []

for fp, c in comenzi:
    op = c.get("op")
    try:
        if op == "add_articol":
            r = ws.max_row + 1
            tpl = find_template_row(c["categorie"], c["tip"])
            copy_row_style(tpl, r)
            data_v = datetime.datetime.strptime(c["data"], "%Y-%m-%d") if c.get("data") else datetime.datetime.now()
            week_prev = ws.cell(row=tpl, column=1).value
            week = (week_prev - 1) if isinstance(week_prev, int) else -900
            etichete = c.get("etichete", "")
            ws.cell(row=r, column=1, value=week)
            ws.cell(row=r, column=2, value=data_v)
            ws.cell(row=r, column=3, value=c["categorie"])
            ws.cell(row=r, column=4, value=c["tip"])
            ws.cell(row=r, column=5, value=c["titlu"])
            ws.cell(row=r, column=6, value=c.get("sursa") or "—")
            ws.cell(row=r, column=7, value=c.get("url") or "")
            ws.cell(row=r, column=9, value=c.get("sponsorizat") or "Nu")
            ws.cell(row=r, column=10, value=etichete or None)
            mapping = slug_row_map()
            slug = [s for s, rr in mapping.items() if rr == r][0]
            write_continut(slug, c.get("continut", ""))
            sync_status(slug, r)
            log.append(f"[add_articol] {slug} (rând {r})")

        elif op == "edit_articol":
            mapping = slug_row_map()
            r = mapping.get(c["slug"])
            if not r:
                log.append(f"[edit_articol] IGNORAT — slug negăsit: {c['slug']}")
                continue
            old_slug = c["slug"]
            if c.get("data"):
                try:
                    ws.cell(row=r, column=2, value=datetime.datetime.strptime(c["data"], "%Y-%m-%d"))
                except ValueError:
                    pass
            ws.cell(row=r, column=3, value=c["categorie"])
            ws.cell(row=r, column=4, value=c["tip"])
            ws.cell(row=r, column=5, value=c["titlu"])
            ws.cell(row=r, column=6, value=c.get("sursa") or "—")
            ws.cell(row=r, column=7, value=c.get("url") or "")
            ws.cell(row=r, column=9, value=c.get("sponsorizat") or "Nu")
            ws.cell(row=r, column=10, value=c.get("etichete") or None)
            new_slug = slugify(c["titlu"])
            # dedup dacă noul slug se ciocnește cu altul existent (rar, dar posibil)
            mapping_after = {s: rr for s, rr in slug_row_map().items() if rr != r}
            base = new_slug
            n = 2
            while new_slug in mapping_after:
                new_slug = f"{base}-{n}"
                n += 1
            rename_article_assets(old_slug, new_slug)
            write_continut(new_slug, c.get("continut", ""))
            sync_status(new_slug, r)
            log.append(f"[edit_articol] {old_slug} -> {new_slug} (rând {r})")

        elif op == "delete_articol":
            pending_delete_articol_slugs.add(c["slug"])

        elif op == "add_intrebare":
            r = aqa_ws.max_row + 1
            for col in range(1, 4):
                src = aqa_ws.cell(row=aqa_ws.max_row, column=col)
                dst = aqa_ws.cell(row=r, column=col)
                dst.font = Font(name=src.font.name, size=src.font.size, bold=src.font.bold, italic=src.font.italic)
                dst.alignment = Alignment(wrapText=True, vertical="center")
            aqa_ws.cell(row=r, column=1, value=c["slug"])
            aqa_ws.cell(row=r, column=2, value=c["intrebare"])
            aqa_ws.cell(row=r, column=3, value=c["raspuns"])
            log.append(f"[add_intrebare] {c['slug']} (rând {r})")

        elif op == "edit_intrebare":
            r = c["row"]
            aqa_ws.cell(row=r, column=1, value=c["slug"])
            aqa_ws.cell(row=r, column=2, value=c["intrebare"])
            aqa_ws.cell(row=r, column=3, value=c["raspuns"])
            log.append(f"[edit_intrebare] rând {r}")

        elif op == "delete_intrebare":
            pending_delete_intrebare_rows.add(c["row"])

        elif op == "rename_sursa":
            n = 0
            for r in range(2, ws.max_row + 1):
                if (ws.cell(row=r, column=6).value or "").strip() == c["veche"]:
                    ws.cell(row=r, column=6, value=c["noua"])
                    n += 1
            log.append(f"[rename_sursa] '{c['veche']}' -> '{c['noua']}' ({n} articole)")

        elif op == "add_categorie":
            slug = slugify(c["nume"])
            if any(cat["nume"] == c["nume"] for cat in categorii):
                log.append(f"[add_categorie] IGNORAT — există deja: {c['nume']}")
                continue
            categorii.append({
                "nume": c["nume"], "slug": slug, "culoare": c.get("culoare") or "#12213A",
                "descriere": c.get("descriere") or "", "exclusaHomepage": bool(c.get("exclusaHomepage")),
            })
            log.append(f"[add_categorie] {c['nume']}")

        elif op == "edit_categorie":
            old = c.get("numeVechi")
            cat = next((x for x in categorii if x["nume"] == old), None)
            if not cat:
                log.append(f"[edit_categorie] IGNORAT — negăsită: {old}")
                continue
            if cat["nume"] != c["nume"]:
                for r in range(2, ws.max_row + 1):
                    if ws.cell(row=r, column=3).value == cat["nume"]:
                        ws.cell(row=r, column=3, value=c["nume"])
                cat["slug"] = slugify(c["nume"])
            cat["nume"] = c["nume"]
            cat["culoare"] = c.get("culoare") or cat["culoare"]
            cat["descriere"] = c.get("descriere") or cat["descriere"]
            cat["exclusaHomepage"] = bool(c.get("exclusaHomepage"))
            log.append(f"[edit_categorie] {old} -> {c['nume']}")

        elif op == "delete_categorie":
            in_use = any(ws.cell(row=r, column=3).value == c["nume"] for r in range(2, ws.max_row + 1))
            if in_use:
                log.append(f"[delete_categorie] REFUZAT — există articole publicate în '{c['nume']}'. Mută-le mai întâi.")
                continue
            categorii[:] = [x for x in categorii if x["nume"] != c["nume"]]
            log.append(f"[delete_categorie] {c['nume']}")

        else:
            log.append(f"OP NECUNOSCUT: {op}")
    except Exception as e:
        log.append(f"EROARE la {op} ({fp}): {e}")

# ---------- ștergeri (rânduri), aplicate ultimele, în ordine descrescătoare ca să nu strice indecșii ----------
if pending_delete_articol_slugs:
    mapping = slug_row_map()
    for slug in pending_delete_articol_slugs:
        for r in range(2, aqa_ws.max_row + 1):
            if (aqa_ws.cell(row=r, column=1).value or "").strip() == slug:
                pending_delete_intrebare_rows.add(r)

for r in sorted(pending_delete_intrebare_rows, reverse=True):
    aqa_ws.delete_rows(r, 1)
    log.append(f"[delete_intrebare] rând {r} șters")

if pending_delete_articol_slugs:
    mapping = slug_row_map()
    rows_to_delete = sorted((mapping[s] for s in pending_delete_articol_slugs if s in mapping), reverse=True)
    for r in rows_to_delete:
        pass
    for slug in pending_delete_articol_slugs:
        r = mapping.get(slug)
        if not r:
            log.append(f"[delete_articol] IGNORAT — slug negăsit: {slug}")
    for r in rows_to_delete:
        ws.delete_rows(r, 1)
    for slug in pending_delete_articol_slugs:
        if slug in mapping:
            txt_path = f"{ARTICLE_TEXT_DIR}/{slug}.txt"
            if os.path.exists(txt_path):
                os.remove(txt_path)
            img_path = f"{IMG_ART_DIR}/{slug}.jpg"
            if os.path.exists(img_path):
                os.remove(img_path)
            img_manifest.pop(slug, None)
            log.append(f"[delete_articol] {slug} șters (conținut + poză + întrebări)")

wb.save(XLSX)
with open(CATEGORII_PATH, "w", encoding="utf-8") as f:
    json.dump(categorii, f, ensure_ascii=False, indent=2)
with open(IMG_MANIFEST_PATH, "w", encoding="utf-8") as f:
    json.dump(img_manifest, f, ensure_ascii=False, indent=2)

for fp, _ in comenzi:
    shutil.move(fp, f"{COMENZI_DONE_DIR}/{os.path.basename(fp)}")

print(f"Procesate {len(comenzi)} comenzi:")
for line in log:
    print(" -", line)

print("\nRulează acum py scripts\\build_site.py și py scripts\\check_links.py ca schimbările să apară pe site.")
