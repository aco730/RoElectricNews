"""Server local pentru admin/ — CRUD instant (fără fișiere de comandă, fără conectare de foldere)
și generare complet automată de articole noi dintr-un link, folosind un API Claude.

Pornire: dublu-click pe pornește-admin.bat (la rădăcina proiectului), sau: py scripts\\admin_server.py
Se deschide automat în browser la http://127.0.0.1:5151/admin/index.html

Cerință: un fișier `.env` la rădăcina proiectului, cu ANTHROPIC_API_KEY=...
(niciodată în cod/git) — necesar doar pentru generarea de articole noi, nu pentru CRUD.
"""
import os, sys, json, datetime, subprocess, threading, webbrowser, time, re
import requests
from bs4 import BeautifulSoup
from flask import Flask, request, jsonify, send_from_directory, redirect
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from fetch_source_images import find_og_image, download_image, HEADERS as FETCH_HEADERS  # noqa: E402
from fetch_source_dates import find_published_date  # noqa: E402
from gen_article_images import make_article_image, CAT_SLUG as GEN_CAT_SLUG  # noqa: E402


def load_env(path):
    if os.path.exists(path):
        for line in open(path, encoding="utf-8"):
            line = line.strip()
            if line and not line.startswith("#") and "=" in line:
                k, v = line.split("=", 1)
                os.environ.setdefault(k.strip(), v.strip())


load_env(f"{ROOT}/.env")

XLSX = f"{ROOT}/electric-news-calendar-editorial.xlsx"
CATEGORII_PATH = f"{ROOT}/categorii.json"
ARTICLE_TEXT_DIR = f"{ROOT}/articole-text"
IMG_ART_DIR = f"{ROOT}/site-local/assets/images/articole"
IMG_MANIFEST_PATH = f"{IMG_ART_DIR}/.image-manifest.json"

app = Flask(__name__)


def slugify(s):
    repl = {"ă":"a","â":"a","î":"i","ș":"s","ş":"s","ț":"t","ţ":"t","Ă":"a","Â":"a","Î":"i","Ș":"s","Ț":"t"}
    out = "".join(repl.get(ch, ch) for ch in s)
    out = out.lower()
    keep = [ch if ch.isalnum() else (" " if ch in " -_/" else "") for ch in out]
    out = "".join(keep)
    return "-".join(out.split())[:80].strip("-")


def load_wb():
    return openpyxl.load_workbook(XLSX)


def load_categorii():
    return json.load(open(CATEGORII_PATH, encoding="utf-8"))


def save_categorii(cats):
    with open(CATEGORII_PATH, "w", encoding="utf-8") as f:
        json.dump(cats, f, ensure_ascii=False, indent=2)


def load_img_manifest():
    return json.load(open(IMG_MANIFEST_PATH, encoding="utf-8")) if os.path.exists(IMG_MANIFEST_PATH) else {}


def save_img_manifest(m):
    with open(IMG_MANIFEST_PATH, "w", encoding="utf-8") as f:
        json.dump(m, f, ensure_ascii=False, indent=2)


def slug_row_map(ws):
    mapping, seen = {}, {}
    for r in range(2, ws.max_row + 1):
        titlu = ws.cell(row=r, column=5).value
        if not titlu:
            continue
        base = slugify(titlu)
        slug = base
        n = 2
        while slug in seen:
            slug = f"{base}-{n}"
            n += 1
        seen[slug] = True
        mapping[slug] = r
    return mapping


def find_template_row(ws, categorie, tip):
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=3).value == categorie and ws.cell(row=r, column=4).value == tip:
            return r
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=4).value == tip:
            return r
    return 2


def copy_row_style(ws, src_row, dst_row):
    for col in range(1, 11):
        src_cell = ws.cell(row=src_row, column=col)
        dst_cell = ws.cell(row=dst_row, column=col)
        dst_cell.font = Font(**{k: getattr(src_cell.font, k) for k in ["name", "size", "bold", "italic", "color"]})
        if src_cell.fill and src_cell.fill.fgColor and src_cell.fill.fgColor.rgb not in (None, "00000000"):
            dst_cell.fill = PatternFill(fill_type="solid", fgColor=src_cell.fill.fgColor.rgb)
        dst_cell.alignment = Alignment(**{k: getattr(src_cell.alignment, k) for k in ["horizontal", "vertical", "wrapText"]})
    ws.cell(row=dst_row, column=2).number_format = ws.cell(row=src_row, column=2).number_format


def rebuild_site():
    env = dict(os.environ, PYTHONIOENCODING="utf-8")
    build = subprocess.run([sys.executable, f"{ROOT}/scripts/build_site.py"], capture_output=True,
                            text=True, encoding="utf-8", errors="replace", cwd=ROOT, env=env)
    check = subprocess.run([sys.executable, f"{ROOT}/scripts/check_links.py"], capture_output=True,
                            text=True, encoding="utf-8", errors="replace", cwd=ROOT, env=env)
    ok = "niciun link mort" in check.stdout
    return {
        "build_ok": build.returncode == 0,
        "build_out": build.stdout[-1500:],
        "check_ok": ok,
        "check_out": check.stdout[-1500:],
    }


# ---------- CRUD articole / întrebări / surse / categorii ----------

@app.route("/api/articol/add", methods=["POST"])
def api_add_articol():
    c = request.json
    wb = load_wb()
    ws = wb["Calendar editorial"]
    r = ws.max_row + 1
    tpl = find_template_row(ws, c["categorie"], c["tip"])
    copy_row_style(ws, tpl, r)
    try:
        data_v = datetime.datetime.strptime(c["data"], "%Y-%m-%d") if c.get("data") else datetime.datetime.now()
    except ValueError:
        data_v = datetime.datetime.now()
    week_prev = ws.cell(row=tpl, column=1).value
    week = (week_prev - 1) if isinstance(week_prev, int) else -900
    ws.cell(row=r, column=1, value=week)
    ws.cell(row=r, column=2, value=data_v)
    ws.cell(row=r, column=3, value=c["categorie"])
    ws.cell(row=r, column=4, value=c["tip"])
    ws.cell(row=r, column=5, value=c["titlu"])
    ws.cell(row=r, column=6, value=c.get("sursa") or "—")
    ws.cell(row=r, column=7, value=c.get("url") or "")
    ws.cell(row=r, column=9, value=c.get("sponsorizat") or "Nu")
    ws.cell(row=r, column=10, value=c.get("etichete") or None)
    mapping = slug_row_map(ws)
    slug = next(s for s, rr in mapping.items() if rr == r)
    continut = c.get("continut", "")
    if continut.strip():
        with open(f"{ARTICLE_TEXT_DIR}/{slug}.txt", "w", encoding="utf-8") as f:
            f.write(continut.strip() + "\n")
    ws.cell(row=r, column=8, value="Publicat" if os.path.exists(f"{ARTICLE_TEXT_DIR}/{slug}.txt") else "În așteptare")
    wb.save(XLSX)
    result = rebuild_site()
    return jsonify({"ok": True, "slug": slug, "rebuild": result})


@app.route("/api/articol/edit", methods=["POST"])
def api_edit_articol():
    c = request.json
    wb = load_wb()
    ws = wb["Calendar editorial"]
    mapping = slug_row_map(ws)
    r = mapping.get(c["slug"])
    if not r:
        return jsonify({"ok": False, "error": f"Articolul '{c['slug']}' nu a fost găsit."}), 404
    old_slug = c["slug"]
    if c.get("data"):
        try:
            ws.cell(row=r, column=2, value=datetime.datetime.strptime(c["data"], "%Y-%m-%d"))
        except ValueError:
            pass
    ws.cell(row=r, column=3, value=c["categorie"])
    ws.cell(row=r, column=4, value=c["tip"])
    ws.cell(row=r, column=5, value=c["titlu"])
    ws.cell(row=r, column=6, value=c.get("sursa") or "—")
    ws.cell(row=r, column=7, value=c.get("url") or "")
    ws.cell(row=r, column=9, value=c.get("sponsorizat") or "Nu")
    ws.cell(row=r, column=10, value=c.get("etichete") or None)

    new_slug = slugify(c["titlu"])
    mapping_after = {s: rr for s, rr in slug_row_map(ws).items() if rr != r}
    base, n = new_slug, 2
    while new_slug in mapping_after:
        new_slug = f"{base}-{n}"
        n += 1

    if old_slug != new_slug:
        old_txt = f"{ARTICLE_TEXT_DIR}/{old_slug}.txt"
        if os.path.exists(old_txt):
            os.replace(old_txt, f"{ARTICLE_TEXT_DIR}/{new_slug}.txt")
        old_img = f"{IMG_ART_DIR}/{old_slug}.jpg"
        if os.path.exists(old_img):
            os.replace(old_img, f"{IMG_ART_DIR}/{new_slug}.jpg")
        manifest = load_img_manifest()
        if old_slug in manifest:
            manifest[new_slug] = manifest.pop(old_slug)
            save_img_manifest(manifest)
        aqa_ws = wb["Q&A Articole"]
        for rr in range(2, aqa_ws.max_row + 1):
            if (aqa_ws.cell(row=rr, column=1).value or "").strip() == old_slug:
                aqa_ws.cell(row=rr, column=1, value=new_slug)

    continut = c.get("continut", "")
    if continut.strip():
        with open(f"{ARTICLE_TEXT_DIR}/{new_slug}.txt", "w", encoding="utf-8") as f:
            f.write(continut.strip() + "\n")
    ws.cell(row=r, column=8, value="Publicat" if os.path.exists(f"{ARTICLE_TEXT_DIR}/{new_slug}.txt") else "În așteptare")
    wb.save(XLSX)
    result = rebuild_site()
    return jsonify({"ok": True, "slug": new_slug, "rebuild": result})


@app.route("/api/articol/delete", methods=["POST"])
def api_delete_articol():
    slug = request.json["slug"]
    wb = load_wb()
    ws = wb["Calendar editorial"]
    aqa_ws = wb["Q&A Articole"]
    mapping = slug_row_map(ws)
    r = mapping.get(slug)
    if not r:
        return jsonify({"ok": False, "error": f"Articolul '{slug}' nu a fost găsit."}), 404
    qa_rows = [rr for rr in range(2, aqa_ws.max_row + 1) if (aqa_ws.cell(row=rr, column=1).value or "").strip() == slug]
    for rr in sorted(qa_rows, reverse=True):
        aqa_ws.delete_rows(rr, 1)
    ws.delete_rows(r, 1)
    for path in (f"{ARTICLE_TEXT_DIR}/{slug}.txt", f"{IMG_ART_DIR}/{slug}.jpg"):
        if os.path.exists(path):
            os.remove(path)
    manifest = load_img_manifest()
    if manifest.pop(slug, None) is not None:
        save_img_manifest(manifest)
    wb.save(XLSX)
    result = rebuild_site()
    return jsonify({"ok": True, "rebuild": result})


@app.route("/api/intrebare/add", methods=["POST"])
def api_add_intrebare():
    c = request.json
    wb = load_wb()
    aqa_ws = wb["Q&A Articole"]
    r = aqa_ws.max_row + 1
    for col in range(1, 4):
        src = aqa_ws.cell(row=aqa_ws.max_row, column=col)
        dst = aqa_ws.cell(row=r, column=col)
        dst.font = Font(name=src.font.name, size=src.font.size, bold=src.font.bold, italic=src.font.italic)
        dst.alignment = Alignment(wrapText=True, vertical="center")
    aqa_ws.cell(row=r, column=1, value=c["slug"])
    aqa_ws.cell(row=r, column=2, value=c["intrebare"])
    aqa_ws.cell(row=r, column=3, value=c["raspuns"])
    wb.save(XLSX)
    result = rebuild_site()
    return jsonify({"ok": True, "rebuild": result})


@app.route("/api/intrebare/edit", methods=["POST"])
def api_edit_intrebare():
    c = request.json
    wb = load_wb()
    aqa_ws = wb["Q&A Articole"]
    r = c["row"]
    aqa_ws.cell(row=r, column=1, value=c["slug"])
    aqa_ws.cell(row=r, column=2, value=c["intrebare"])
    aqa_ws.cell(row=r, column=3, value=c["raspuns"])
    wb.save(XLSX)
    result = rebuild_site()
    return jsonify({"ok": True, "rebuild": result})


@app.route("/api/intrebare/delete", methods=["POST"])
def api_delete_intrebare():
    r = request.json["row"]
    wb = load_wb()
    aqa_ws = wb["Q&A Articole"]
    aqa_ws.delete_rows(r, 1)
    wb.save(XLSX)
    result = rebuild_site()
    return jsonify({"ok": True, "rebuild": result})


@app.route("/api/sursa/rename", methods=["POST"])
def api_rename_sursa():
    c = request.json
    wb = load_wb()
    ws = wb["Calendar editorial"]
    n = 0
    for r in range(2, ws.max_row + 1):
        if (ws.cell(row=r, column=6).value or "").strip() == c["veche"]:
            ws.cell(row=r, column=6, value=c["noua"])
            n += 1
    wb.save(XLSX)
    result = rebuild_site()
    return jsonify({"ok": True, "count": n, "rebuild": result})


@app.route("/api/categorie/add", methods=["POST"])
def api_add_categorie():
    c = request.json
    cats = load_categorii()
    if any(x["nume"] == c["nume"] for x in cats):
        return jsonify({"ok": False, "error": "Există deja o categorie cu acest nume."}), 400
    cats.append({
        "nume": c["nume"], "slug": slugify(c["nume"]), "culoare": c.get("culoare") or "#12213A",
        "descriere": c.get("descriere") or "", "exclusaHomepage": bool(c.get("exclusaHomepage")),
    })
    save_categorii(cats)
    result = rebuild_site()
    warn = "Categorie nouă — dacă publici articole în ea, adaugă manual o poză de fundal la site-local/assets/images/categorii/<slug>.jpg."
    return jsonify({"ok": True, "warning": warn, "rebuild": result})


@app.route("/api/categorie/edit", methods=["POST"])
def api_edit_categorie():
    c = request.json
    cats = load_categorii()
    cat = next((x for x in cats if x["nume"] == c["numeVechi"]), None)
    if not cat:
        return jsonify({"ok": False, "error": "Categorie negăsită."}), 404
    if cat["nume"] != c["nume"]:
        wb = load_wb()
        ws = wb["Calendar editorial"]
        for r in range(2, ws.max_row + 1):
            if ws.cell(row=r, column=3).value == cat["nume"]:
                ws.cell(row=r, column=3, value=c["nume"])
        wb.save(XLSX)
        cat["slug"] = slugify(c["nume"])
    cat["nume"] = c["nume"]
    cat["culoare"] = c.get("culoare") or cat["culoare"]
    cat["descriere"] = c.get("descriere") or cat["descriere"]
    cat["exclusaHomepage"] = bool(c.get("exclusaHomepage"))
    save_categorii(cats)
    result = rebuild_site()
    return jsonify({"ok": True, "rebuild": result})


@app.route("/api/categorie/delete", methods=["POST"])
def api_delete_categorie():
    nume = request.json["nume"]
    wb = load_wb()
    ws = wb["Calendar editorial"]
    in_use = any(ws.cell(row=r, column=3).value == nume for r in range(2, ws.max_row + 1))
    if in_use:
        return jsonify({"ok": False, "error": f"Există articole publicate în '{nume}' — mută-le sau șterge-le mai întâi."}), 400
    cats = load_categorii()
    cats[:] = [x for x in cats if x["nume"] != nume]
    save_categorii(cats)
    result = rebuild_site()
    return jsonify({"ok": True, "rebuild": result})


# ---------- generare automată de articole noi (Claude API) ----------

def call_claude(prompt, max_tokens=3000):
    api_key = os.environ.get("ANTHROPIC_API_KEY")
    if not api_key:
        raise RuntimeError("Lipsește ANTHROPIC_API_KEY din .env la rădăcina proiectului.")
    resp = requests.post(
        "https://api.anthropic.com/v1/messages",
        headers={"x-api-key": api_key, "anthropic-version": "2023-06-01", "content-type": "application/json"},
        json={"model": "claude-sonnet-5", "max_tokens": max_tokens, "messages": [{"role": "user", "content": prompt}]},
        timeout=90,
    )
    resp.raise_for_status()
    data = resp.json()
    return "".join(block.get("text", "") for block in data.get("content", []))


def extract_json(text):
    text = text.strip()
    if text.startswith("```"):
        text = re.sub(r"^```[a-z]*\n?", "", text)
        text = re.sub(r"\n?```$", "", text)
    return json.loads(text)


CATEGORII_NUME = [c["nume"] for c in load_categorii()]

SETARI_PATH = f"{ROOT}/setari.json"
SETARI_DEFAULT = {
    "ton": "tehnic-sobru",
    "lungime_min": 350,
    "lungime_max": 500,
    "numar_intrebari": 5,
    "focus_extractie": "cifre exacte, citate directe (cu nume și funcție dacă există), termene și date concrete, context legislativ",
    "poze_strategie": "og_pexels_generata",
    "poze_cuvinte_cheie_extra": "",
}
TON_DESCRIERI = {
    "tehnic-sobru": "ton tehnic și sobru, NU promoțional",
    "profesional": "ton profesional, formal, orientat spre precizie și claritate",
    "informativ-simplu": "ton informativ, accesibil, propoziții scurte, fără jargon inutil",
    "conversational": "ton conversațional, prietenos, ca și cum ai explica unui prieten",
}


def load_setari():
    if os.path.exists(SETARI_PATH):
        s = dict(SETARI_DEFAULT)
        s.update(json.load(open(SETARI_PATH, encoding="utf-8")))
        return s
    return dict(SETARI_DEFAULT)


def save_setari(s):
    with open(SETARI_PATH, "w", encoding="utf-8") as f:
        json.dump(s, f, ensure_ascii=False, indent=2)


PROMPT_LINK = """Ești un redactor pentru Electric NEWS, un site românesc de nișă despre electric, smart home, fotovoltaic și baterii de stocare.
Scrie un articol REAL bazat STRICT pe textul sursă de mai jos — niciodată nu inventezi cifre, citate sau fapte care nu apar în text.
Dacă un detaliu așteptat lipsește din sursă, spune asta direct în articol, nu completa cu presupuneri.
La extragerea faptelor, acordă prioritate la: {focus_extractie}.

Categorii disponibile (alege exact una, copiază exact numele): {categorii}
{hints}
Text sursă (extras automat de pe {url}):
\"\"\"
{page_text}
\"\"\"

Răspunde DOAR cu un obiect JSON valid, fără text explicativ, fără code fences, cu exact aceste chei:
{{
  "titlu": "titlu editorial concret, în română, nu traducere mecanică",
  "categorie": "una dintre categoriile de mai sus, exact",
  "sursa_nume": "nume scurt al publicației/site-ului sursă",
  "articol": "{lungime_min}-{lungime_max} cuvinte, {ton_descriere}, paragrafe separate prin \\n\\n, ultimul paragraf menționează explicit sursa (\\"Sursă: X.\\")",
  "intrebari": [{{"intrebare": "...", "raspuns": "1-2 propoziții, extrase direct din articol"}}],
  "etichete": "0-2 etichete separate prin virgulă, doar din: DIY, Opinie, Prețuri, Tehnologie nouă — sau string gol",
  "pexels_query": "2-4 cuvinte cheie ÎN ENGLEZĂ, descriptive și vizuale, pentru căutarea unei fotografii de stoc potrivite temei articolului (ex. \\"solar panels rooftop\\")"
}}
"intrebari" trebuie să aibă exact {numar_intrebari} elemente."""

PROMPT_IDEE = """Ești un redactor pentru Electric NEWS, un site românesc de nișă despre electric, smart home, fotovoltaic și baterii de stocare.
Scrie un ghid original (~{lungime_min}-{lungime_max_idee} cuvinte) pe baza ideii de mai jos, folosind expertiză tehnică generală a domeniului — fără cifre sau afirmații legale inventate cu falsă precizie; dacă un detaliu tehnic nu e sigur, formulează general ("de regulă", "în practică"). {ton_descriere}.

Categorii disponibile (alege exact una, copiază exact numele): {categorii}

Idee primită: \"\"\"{page_text}\"\"\"

Răspunde DOAR cu un obiect JSON valid, fără text explicativ, fără code fences, cu exact aceste chei:
{{
  "titlu": "titlu editorial concret",
  "categorie": "una dintre categoriile de mai sus, exact",
  "sursa_nume": "—",
  "articol": "{lungime_min}-{lungime_max_idee} cuvinte, paragrafe separate prin \\n\\n",
  "intrebari": [{{"intrebare": "...", "raspuns": "..."}}],
  "etichete": "0-2 etichete separate prin virgulă, doar din: DIY, Opinie, Prețuri, Tehnologie nouă — sau string gol",
  "pexels_query": "2-4 cuvinte cheie ÎN ENGLEZĂ, descriptive și vizuale, pentru căutarea unei fotografii de stoc potrivite temei articolului"
}}
"intrebari" trebuie să aibă exact {numar_intrebari} elemente."""

PROMPT_YOUTUBE = """Ești un redactor pentru Electric NEWS. Ai primit un transcript brut (posibil cu erori STT) al unui video YouTube relevant pentru electric/smart home/fotovoltaic/baterii.
Scrie un rezumat editorial SCURT (~250-300 cuvinte, mai scurt decât un articol normal), citibil în ~2 minute, bazat STRICT pe conținutul transcriptului. {ton_descriere}.

Categorii disponibile: {categorii}. Pentru conținut YouTube alege "YouTube".

Transcript:
\"\"\"
{page_text}
\"\"\"

Răspunde DOAR cu un obiect JSON valid, fără text explicativ, fără code fences:
{{
  "titlu": "titlu relevant din conținutul video-ului",
  "categorie": "YouTube",
  "sursa_nume": "nume canal, dacă apare, altfel —",
  "articol": "250-300 cuvinte, paragrafe separate prin \\n\\n",
  "intrebari": [{{"intrebare": "...", "raspuns": "..."}}],
  "etichete": ""
}}
"intrebari" trebuie să aibă exact {numar_intrebari} elemente."""


def fetch_page(url):
    r = requests.get(url, headers=FETCH_HEADERS, timeout=20)
    r.raise_for_status()
    soup = BeautifulSoup(r.text, "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    text = "\n".join(l.strip() for l in soup.get_text("\n").splitlines() if l.strip())
    return text[:9000], r.text


def try_pexels_image(slug, query):
    api_key = os.environ.get("PEXELS_API_KEY")
    if not api_key or not query:
        return False
    try:
        r = requests.get(
            "https://api.pexels.com/v1/search",
            headers={"Authorization": api_key},
            params={"query": query, "per_page": 5, "orientation": "landscape"},
            timeout=20,
        )
        r.raise_for_status()
        photos = r.json().get("photos", [])
        if not photos:
            return False
        photo = photos[0]
        img_resp = requests.get(photo["src"]["large2x"], timeout=30)
        img_resp.raise_for_status()
        with open(f"{IMG_ART_DIR}/{slug}.jpg", "wb") as f:
            f.write(img_resp.content)
        manifest = load_img_manifest()
        manifest[slug] = {
            "source": "stock-pexels", "photographer": photo.get("photographer"),
            "photographer_url": photo.get("photographer_url"), "pexels_url": photo.get("url"), "query": query,
            "fetched_at": datetime.datetime.now().isoformat(timespec="seconds"),
        }
        save_img_manifest(manifest)
        return True
    except Exception:
        return False


def try_og_image(slug, url, raw_html):
    if not url or not raw_html:
        return False
    img_url = find_og_image(url, raw_html)
    if not img_url:
        return False
    try:
        img = download_image(img_url)
        img.save(f"{IMG_ART_DIR}/{slug}.jpg", quality=90)
        manifest = load_img_manifest()
        manifest[slug] = {"source": "external", "origin_image_url": img_url, "page_url": url,
                           "fetched_at": datetime.datetime.now().isoformat(timespec="seconds")}
        save_img_manifest(manifest)
        return True
    except Exception:
        return False


def acquire_image(slug, categorie, titlu, url, raw_html, pexels_query_hint, setari):
    strategie = setari.get("poze_strategie", "og_pexels_generata")
    steps = {
        "og_pexels_generata": ["og", "pexels", "generata"],
        "og_generata": ["og", "generata"],
        "pexels_generata": ["pexels", "generata"],
        "doar_generata": ["generata"],
    }.get(strategie, ["og", "pexels", "generata"])

    extra = setari.get("poze_cuvinte_cheie_extra", "").strip()
    query = (pexels_query_hint or categorie).strip()
    if extra:
        query = f"{query} {extra}"

    for step in steps:
        if step == "og" and try_og_image(slug, url, raw_html):
            return
        if step == "pexels" and try_pexels_image(slug, query):
            return
        if step == "generata":
            cat_slug = GEN_CAT_SLUG.get(categorie)
            if cat_slug:
                kind, degrees = make_article_image(slug, cat_slug, titlu, f"{IMG_ART_DIR}/{slug}.jpg")
                manifest = load_img_manifest()
                manifest[slug] = {"source": "generated", "icon": kind, "hue": degrees,
                                   "generated_at": datetime.datetime.now().isoformat(timespec="seconds")}
                save_img_manifest(manifest)
            return


def build_hints(titlu_propus, categorie_sugerata, nota):
    parts = []
    if titlu_propus:
        parts.append(f'Titlu preferat de redacție (folosește-l dacă se potrivește conținutului real, altfel ajustează-l): "{titlu_propus}"')
    if categorie_sugerata:
        parts.append(f'Categorie preferată de redacție (folosește-o dacă se potrivește): "{categorie_sugerata}"')
    if nota:
        parts.append(f"Notă suplimentară de la redacție: {nota}")
    return ("Indicii de la redacție (nu sunt obligatorii dacă nu se potrivesc conținutului real):\n- " + "\n- ".join(parts) + "\n") if parts else ""


def process_one_link(url, sponsorizat, sursa_override=None, titlu_propus="", categorie_sugerata="", nota=""):
    setari = load_setari()
    page_text, raw_html = fetch_page(url)
    hints = build_hints(titlu_propus, categorie_sugerata, nota)
    prompt = PROMPT_LINK.format(
        categorii=", ".join(CATEGORII_NUME), url=url, page_text=page_text, hints=hints,
        lungime_min=setari["lungime_min"], lungime_max=setari["lungime_max"],
        ton_descriere=TON_DESCRIERI.get(setari["ton"], TON_DESCRIERI["tehnic-sobru"]),
        numar_intrebari=setari["numar_intrebari"], focus_extractie=setari["focus_extractie"],
    )
    out = extract_json(call_claude(prompt))

    wb = load_wb()
    ws = wb["Calendar editorial"]
    tpl = find_template_row(ws, out["categorie"], "Agregat")
    r = ws.max_row + 1
    copy_row_style(ws, tpl, r)
    date_v = None
    d, _src = find_published_date(url, raw_html)
    date_v = datetime.datetime(d.year, d.month, d.day) if d else datetime.datetime.now()
    week_prev = ws.cell(row=tpl, column=1).value
    week = (week_prev - 1) if isinstance(week_prev, int) else -900
    ws.cell(row=r, column=1, value=week)
    ws.cell(row=r, column=2, value=date_v)
    ws.cell(row=r, column=3, value=out["categorie"])
    ws.cell(row=r, column=4, value="Agregat")
    ws.cell(row=r, column=5, value=out["titlu"])
    ws.cell(row=r, column=6, value=sursa_override or out.get("sursa_nume") or "—")
    ws.cell(row=r, column=7, value=url)
    ws.cell(row=r, column=8, value="Publicat")
    ws.cell(row=r, column=9, value=sponsorizat)
    ws.cell(row=r, column=10, value=out.get("etichete") or None)

    mapping = slug_row_map(ws)
    slug = next(s for s, rr in mapping.items() if rr == r)
    with open(f"{ARTICLE_TEXT_DIR}/{slug}.txt", "w", encoding="utf-8") as f:
        f.write(out["titlu"] + "\n\n" + out["articol"].strip() + "\n")

    aqa_ws = wb["Q&A Articole"]
    for q in out.get("intrebari", [])[: setari["numar_intrebari"]]:
        rr = aqa_ws.max_row + 1
        aqa_ws.cell(row=rr, column=1, value=slug)
        aqa_ws.cell(row=rr, column=2, value=q["intrebare"])
        aqa_ws.cell(row=rr, column=3, value=q["raspuns"])

    wb.save(XLSX)

    acquire_image(slug, out["categorie"], out["titlu"], url, raw_html, out.get("pexels_query"), setari)

    return slug, out["titlu"], out["categorie"]


CERERI_DIR = f"{ROOT}/cereri-publicare"


@app.route("/api/cerere/salveaza", methods=["POST"])
def api_salveaza_cerere():
    """Salvează cererea în cereri-publicare/ pentru procesare manuală de Claude (fără cost API,
    fără cheie separată) — folosit când generarea automată e dezactivată (vezi discuția despre
    facturare separată Console vs. Claude Pro)."""
    c = request.json
    os.makedirs(CERERI_DIR, exist_ok=True)
    payload = {
        "tip": c["tip"],
        "continut": (c.get("continut") or "").strip(),
        "titlu_propus": (c.get("titlu_propus") or "").strip(),
        "categorie_sugerata": c.get("categorie_sugerata") or "",
        "nota": (c.get("nota") or "").strip(),
        "sponsorizat": c.get("sponsorizat") or "Nu",
        "creat_la": datetime.datetime.now().isoformat(),
    }
    ts = payload["creat_la"].replace(":", "-").replace(".", "-")
    filename = f"cerere-{payload['tip']}-{ts}.json"
    with open(f"{CERERI_DIR}/{filename}", "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
    return jsonify({"ok": True, "filename": filename})


@app.route("/api/genereaza-articol", methods=["POST"])
def api_genereaza_articol():
    c = request.json
    tip = c["tip"]
    sponsorizat = c.get("sponsorizat") or "Nu"
    continut = (c.get("continut") or "").strip()
    generate = []
    errors = []

    try:
        if tip == "link":
            generate = [continut]
        elif tip == "lista":
            generate = [l.strip() for l in continut.splitlines() if l.strip()]
        elif tip == "idee":
            setari = load_setari()
            prompt = PROMPT_IDEE.format(
                categorii=", ".join(CATEGORII_NUME), page_text=continut,
                lungime_min=setari["lungime_min"], lungime_max_idee=setari["lungime_max"] + 50,
                ton_descriere=TON_DESCRIERI.get(setari["ton"], TON_DESCRIERI["tehnic-sobru"]),
                numar_intrebari=setari["numar_intrebari"],
            )
            out = extract_json(call_claude(prompt))
            wb = load_wb()
            ws = wb["Calendar editorial"]
            tpl = find_template_row(ws, out["categorie"], "Propriu")
            r = ws.max_row + 1
            copy_row_style(ws, tpl, r)
            week_prev = ws.cell(row=tpl, column=1).value
            week = (week_prev - 1) if isinstance(week_prev, int) else -900
            ws.cell(row=r, column=1, value=week)
            ws.cell(row=r, column=2, value=datetime.datetime.now())
            ws.cell(row=r, column=3, value=out["categorie"])
            ws.cell(row=r, column=4, value="Propriu")
            ws.cell(row=r, column=5, value=out["titlu"])
            ws.cell(row=r, column=6, value="—")
            ws.cell(row=r, column=7, value="Conținut propriu — fără sursă externă unică")
            ws.cell(row=r, column=8, value="Publicat")
            ws.cell(row=r, column=9, value=sponsorizat)
            ws.cell(row=r, column=10, value=out.get("etichete") or None)
            mapping = slug_row_map(ws)
            slug = next(s for s, rr in mapping.items() if rr == r)
            with open(f"{ARTICLE_TEXT_DIR}/{slug}.txt", "w", encoding="utf-8") as f:
                f.write(out["titlu"] + "\n\n" + out["articol"].strip() + "\n")
            aqa_ws = wb["Q&A Articole"]
            for q in out.get("intrebari", [])[: setari["numar_intrebari"]]:
                rr = aqa_ws.max_row + 1
                aqa_ws.cell(row=rr, column=1, value=slug)
                aqa_ws.cell(row=rr, column=2, value=q["intrebare"])
                aqa_ws.cell(row=rr, column=3, value=q["raspuns"])
            wb.save(XLSX)
            acquire_image(slug, out["categorie"], out["titlu"], None, None, out.get("pexels_query"), setari)
            result = rebuild_site()
            return jsonify({"ok": True, "publicate": [{"slug": slug, "titlu": out["titlu"], "categorie": out["categorie"]}], "erori": [], "rebuild": result})
        elif tip == "youtube":
            setari = load_setari()
            lines = continut.splitlines()
            url = lines[0].strip() if lines and lines[0].strip().startswith("http") else ""
            transcript = "\n".join(lines[1:] if url else lines).strip()
            prompt = PROMPT_YOUTUBE.format(
                categorii=", ".join(CATEGORII_NUME), page_text=transcript[:9000],
                ton_descriere=TON_DESCRIERI.get(setari["ton"], TON_DESCRIERI["tehnic-sobru"]),
                numar_intrebari=setari["numar_intrebari"],
            )
            out = extract_json(call_claude(prompt))
            wb = load_wb()
            ws = wb["Calendar editorial"]
            tpl = find_template_row(ws, "YouTube", "Agregat")
            r = ws.max_row + 1
            copy_row_style(ws, tpl, r)
            week_prev = ws.cell(row=tpl, column=1).value
            week = (week_prev - 1) if isinstance(week_prev, int) else -900
            ws.cell(row=r, column=1, value=week)
            ws.cell(row=r, column=2, value=datetime.datetime.now())
            ws.cell(row=r, column=3, value="YouTube")
            ws.cell(row=r, column=4, value="Agregat")
            ws.cell(row=r, column=5, value=out["titlu"])
            ws.cell(row=r, column=6, value=out.get("sursa_nume") or "—")
            ws.cell(row=r, column=7, value=url)
            ws.cell(row=r, column=8, value="Publicat")
            ws.cell(row=r, column=9, value=sponsorizat)
            ws.cell(row=r, column=10, value=None)
            mapping = slug_row_map(ws)
            slug = next(s for s, rr in mapping.items() if rr == r)
            with open(f"{ARTICLE_TEXT_DIR}/{slug}.txt", "w", encoding="utf-8") as f:
                f.write(out["titlu"] + "\n\n" + out["articol"].strip() + "\n")
            aqa_ws = wb["Q&A Articole"]
            for q in out.get("intrebari", [])[: setari["numar_intrebari"]]:
                rr = aqa_ws.max_row + 1
                aqa_ws.cell(row=rr, column=1, value=slug)
                aqa_ws.cell(row=rr, column=2, value=q["intrebare"])
                aqa_ws.cell(row=rr, column=3, value=q["raspuns"])
            wb.save(XLSX)
            if url:
                try:
                    r2 = requests.get(url, headers=FETCH_HEADERS, timeout=20)
                    if try_og_image(slug, url, r2.text):
                        manifest = load_img_manifest()
                        manifest[slug]["source"] = "youtube-thumbnail"
                        save_img_manifest(manifest)
                    else:
                        acquire_image(slug, "YouTube", out["titlu"], None, None, out.get("pexels_query"), setari)
                except Exception:
                    acquire_image(slug, "YouTube", out["titlu"], None, None, out.get("pexels_query"), setari)
            else:
                acquire_image(slug, "YouTube", out["titlu"], None, None, out.get("pexels_query"), setari)
            result = rebuild_site()
            return jsonify({"ok": True, "publicate": [{"slug": slug, "titlu": out["titlu"], "categorie": "YouTube"}], "erori": [], "rebuild": result})
        else:
            return jsonify({"ok": False, "error": f"Tip necunoscut: {tip}"}), 400
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 500

    titlu_propus = c.get("titlu_propus") or ""
    categorie_sugerata = c.get("categorie_sugerata") or ""
    nota = c.get("nota") or ""
    publicate = []
    for url in generate:
        try:
            slug, titlu, categorie = process_one_link(url, sponsorizat, titlu_propus=titlu_propus, categorie_sugerata=categorie_sugerata, nota=nota)
            publicate.append({"slug": slug, "titlu": titlu, "categorie": categorie})
        except Exception as e:
            errors.append({"url": url, "error": str(e)})

    result = rebuild_site() if publicate else {"build_ok": True, "check_ok": True, "build_out": "", "check_out": ""}
    return jsonify({"ok": True, "publicate": publicate, "erori": errors, "rebuild": result})


# ---------- Modul 4: agregator (propune candidați, nu publică automat) ----------

@app.route("/api/agregator/scaneaza", methods=["POST"])
def api_agregator_scaneaza():
    c = request.json
    url = c["url"].strip()
    try:
        page_text, raw_html = fetch_page(url)
    except Exception as e:
        return jsonify({"ok": False, "error": f"Nu am putut accesa sursa: {e}"}), 502

    NAV_PATH_HINTS = (
        "/tag/", "/tags/", "/category/", "/categorie/", "/author/", "/pagina/", "/page/",
        "/despre", "/contact", "/membri", "/inscriere", "/login", "/cont", "/cos", "/wp-content",
        "/wp-admin", "/feed", "/politica", "/termeni", "/privacy", "/sitemap",
    )

    def looks_like_article(href, text):
        from urllib.parse import urlparse
        parsed = urlparse(href)
        path = parsed.path.rstrip("/")
        if parsed.query:
            return False
        if any(h in path.lower() for h in NAV_PATH_HINTS):
            return False
        # sloguri de articol real: path lung, descriptiv (multe cuvinte separate prin cratimă)
        segment = path.rsplit("/", 1)[-1]
        if segment.count("-") < 3 and len(segment) < 40:
            return False
        return True

    soup = BeautifulSoup(raw_html, "html.parser")
    seen, links = set(), []
    for a in soup.find_all("a", href=True):
        href = a["href"]
        text = a.get_text(strip=True)
        if not text or len(text) < 15:
            continue
        if href.startswith("/"):
            from urllib.parse import urljoin
            href = urljoin(url, href)
        if not href.startswith("http") or href in seen:
            continue
        if not looks_like_article(href, text):
            continue
        seen.add(href)
        links.append({"titlu": text[:160], "url": href})

    wb = load_wb()
    ws = wb["Calendar editorial"]
    existing_urls = {(ws.cell(row=r, column=7).value or "").strip() for r in range(2, ws.max_row + 1)}
    candidati = [l for l in links if l["url"] not in existing_urls][:40]
    return jsonify({"ok": True, "candidati": candidati})


@app.route("/api/agregator/publica", methods=["POST"])
def api_agregator_publica():
    c = request.json
    urls = c.get("urls", [])
    sponsorizat = c.get("sponsorizat") or "Nu"
    publicate, errors = [], []
    for url in urls:
        try:
            slug, titlu, categorie = process_one_link(url, sponsorizat)
            publicate.append({"slug": slug, "titlu": titlu, "categorie": categorie})
        except Exception as e:
            errors.append({"url": url, "error": str(e)})
    result = rebuild_site() if publicate else {"build_ok": True, "check_ok": True, "build_out": "", "check_out": ""}
    return jsonify({"ok": True, "publicate": publicate, "erori": errors, "rebuild": result})


@app.route("/api/status", methods=["GET"])
def api_status():
    return jsonify({"ok": True, "are_cheie_api": bool(os.environ.get("ANTHROPIC_API_KEY"))})


@app.route("/api/stare", methods=["GET"])
def api_stare():
    index_path = f"{ROOT}/site-local/index.html"
    build_time = datetime.datetime.fromtimestamp(os.path.getmtime(index_path)).isoformat(timespec="seconds") if os.path.exists(index_path) else None
    wb = load_wb()
    ws = wb["Calendar editorial"]
    total_articole = sum(1 for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=5).value)
    return jsonify({
        "ok": True,
        "cheie_anthropic": bool(os.environ.get("ANTHROPIC_API_KEY")),
        "cheie_pexels": bool(os.environ.get("PEXELS_API_KEY")),
        "goatcounter_conectat": os.path.exists(f"{ROOT}/goatcounter_site_code.txt"),
        "goatcounter_stats": os.path.exists(f"{ROOT}/goatcounter_stats.json"),
        "ultima_regenerare": build_time,
        "total_articole": total_articole,
    })


@app.route("/api/regenereaza", methods=["POST"])
def api_regenereaza():
    result = rebuild_site()
    return jsonify({"ok": True, "rebuild": result})


@app.route("/api/setari", methods=["GET"])
def api_get_setari():
    return jsonify({"ok": True, "setari": load_setari()})


@app.route("/api/setari", methods=["POST"])
def api_save_setari():
    c = request.json
    setari = load_setari()
    for key in SETARI_DEFAULT:
        if key in c:
            setari[key] = c[key]
    save_setari(setari)
    return jsonify({"ok": True, "setari": setari})


# ---------- fișiere statice (admin/, site-local/, etc.) ----------

@app.route("/")
def root():
    return redirect("/admin/index.html")


@app.route("/<path:path>")
def static_files(path):
    return send_from_directory(ROOT, path)


def open_browser():
    time.sleep(1.0)
    webbrowser.open("http://127.0.0.1:5151/admin/index.html")


if __name__ == "__main__":
    threading.Thread(target=open_browser, daemon=True).start()
    print("Server admin pornit — http://127.0.0.1:5151/admin/index.html")
    print("Lasă această fereastră deschisă cât timp folosești admin-ul. Închide-o (sau Ctrl+C) ca să oprești serverul.")
    app.run(host="127.0.0.1", port=5151, debug=False)
