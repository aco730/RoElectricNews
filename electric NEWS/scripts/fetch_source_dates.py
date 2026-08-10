import os, re, json, sys, datetime
import requests
from bs4 import BeautifulSoup
import openpyxl
from copy import copy

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = f"{ROOT}/electric-news-calendar-editorial.xlsx"

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                  "(KHTML, like Gecko) Chrome/124.0 Safari/537.36",
    "Accept-Language": "ro-RO,ro;q=0.9,en;q=0.8",
}

def parse_date(s):
    s = s.strip()
    for fmt in ("%Y-%m-%dT%H:%M:%S%z", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.datetime.strptime(s[:len(fmt.replace('%z','+0000'))] if '%z' in fmt else s[:len(fmt)], fmt).date()
        except ValueError:
            continue
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return datetime.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    return None

def find_published_date(url, html):
    soup = BeautifulSoup(html, "html.parser")
    for prop in ["article:published_time", "og:article:published_time", "article:modified_time"]:
        tag = soup.find("meta", {"property": prop})
        if tag and tag.get("content"):
            d = parse_date(tag["content"])
            if d:
                return d, f"meta[{prop}]"
    for name in ["date", "publish-date", "publishdate", "sailthru.date", "dc.date", "dc.date.issued"]:
        tag = soup.find("meta", {"name": name})
        if tag and tag.get("content"):
            d = parse_date(tag["content"])
            if d:
                return d, f"meta[name={name}]"
    for script in soup.find_all("script", {"type": "application/ld+json"}):
        try:
            data = json.loads(script.string or "{}")
        except Exception:
            continue
        items = data if isinstance(data, list) else [data]
        for it in items:
            if not isinstance(it, dict):
                continue
            for key in ("datePublished", "dateCreated", "uploadDate"):
                if it.get(key):
                    d = parse_date(str(it[key]))
                    if d:
                        return d, f"ld+json[{key}]"
    time_tag = soup.find("time", {"datetime": True})
    if time_tag:
        d = parse_date(time_tag["datetime"])
        if d:
            return d, "time[datetime]"
    if "youtube.com/watch" in url:
        m = re.search(r'"(?:publishDate|uploadDate)":"(\d{4}-\d{2}-\d{2})', html)
        if m:
            return parse_date(m.group(1)), "youtube-json"
    return None, None

def main(dry_run=True):
    wb = openpyxl.load_workbook(XLSX)
    ws = wb["Calendar editorial"]

    cache = {}
    results = []
    for r in range(2, ws.max_row + 1):
        titlu = ws.cell(row=r, column=5).value
        status = ws.cell(row=r, column=8).value
        url = (ws.cell(row=r, column=7).value or "").strip()
        if not titlu or status != "Publicat" or not url.startswith("http"):
            continue
        if url not in cache:
            try:
                resp = requests.get(url, headers=HEADERS, timeout=20)
                resp.raise_for_status()
                cache[url] = find_published_date(url, resp.text)
            except Exception as e:
                cache[url] = (None, f"eroare: {e}")
        date, source = cache[url]
        old_date = ws.cell(row=r, column=2).value
        results.append((r, titlu[:55], old_date, date, source, url))

    for r, titlu, old_date, date, source, url in results:
        status_str = f"{old_date} -> {date} ({source})" if date else f"{old_date} -> NEGASIT ({source})"
        print(f"{r:3} | {status_str} | {titlu}")

    if not dry_run:
        changed = 0
        for r, titlu, old_date, date, source, url in results:
            if date:
                dt = datetime.datetime(date.year, date.month, date.day)
                cell = ws.cell(row=r, column=2)
                cell.value = dt
                changed += 1
        wb.save(XLSX)
        print(f"\nSalvat: {changed} date actualizate din {len(results)} rânduri cu sursă.")
    else:
        print(f"\n[DRY RUN] {sum(1 for x in results if x[3])} din {len(results)} găsite. Rulează cu --apply ca să salvezi.")

if __name__ == "__main__":
    main(dry_run="--apply" not in sys.argv)
