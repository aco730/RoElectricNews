import openpyxl, os, sys, json, datetime, subprocess
from copy import copy
from openpyxl.styles import PatternFill

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = f"{ROOT}/electric-news-calendar-editorial.xlsx"
TRANSCRIPT_DIR = f"{ROOT}/transcripturi-youtube"
MANIFEST = f"{TRANSCRIPT_DIR}/.processed-manifest.json"

NA_TEXT = "Conținut propriu — fără sursă externă unică"
YOUTUBE_RED = "FFFF0000"

DIACRITICE = {"ă":"a","â":"a","î":"i","ș":"s","ş":"s","ț":"t","ţ":"t"}
def slugify(s):
    out = "".join(DIACRITICE.get(c, c) for c in s.lower())
    keep = [c if c.isalnum() else (" " if c in " -_/" else "") for c in out]
    out = "".join(keep)
    return "-".join(out.split())[:80].strip("-")

def main(transcript_file, titlu, sursa="YouTube"):
    path = f"{TRANSCRIPT_DIR}/{transcript_file}"
    content = open(path, encoding="utf-8").read().strip()
    lines = content.split("\n")
    url = lines[0].strip() if lines and lines[0].strip().startswith("http") else ""

    wb = openpyxl.load_workbook(XLSX)
    ws = wb["Calendar editorial"]
    last_row = ws.max_row
    last_date = ws.cell(row=last_row, column=2).value
    first_date = ws.cell(row=2, column=2).value
    new_date = last_date + datetime.timedelta(days=1)
    new_week = ((new_date - first_date).days // 7) + 1
    new_row = last_row + 1

    ref_row = 2  # Agregat + URL real, stil de referinta
    for col in range(1, 10):
        src = ws.cell(row=ref_row, column=col)
        dst = ws.cell(row=new_row, column=col)
        dst.font, dst.fill = copy(src.font), copy(src.fill)
        dst.border, dst.alignment = copy(src.border), copy(src.alignment)
        dst.number_format = src.number_format

    ws.cell(row=new_row, column=1).value = new_week
    ws.cell(row=new_row, column=2).value = new_date
    ws.cell(row=new_row, column=3).value = "YouTube"
    ws.cell(row=new_row, column=3).fill = PatternFill(fill_type="solid", fgColor=YOUTUBE_RED)
    ws.cell(row=new_row, column=4).value = "Agregat"
    ws.cell(row=new_row, column=5).value = titlu
    ws.cell(row=new_row, column=6).value = sursa
    ws.cell(row=new_row, column=7).value = url or NA_TEXT
    ws.cell(row=new_row, column=8).value = "Publicat"  # continutul e deja scris in articole-text/ inainte sa ruleze acest script
    ws.cell(row=new_row, column=9).value = "Nu"  # Sponsorizat, implicit Nu
    ws.row_dimensions[new_row].height = 34

    wb.save(XLSX)

    manifest = json.load(open(MANIFEST, encoding="utf-8")) if os.path.exists(MANIFEST) else {}
    manifest[transcript_file] = {
        "slug": slugify(titlu), "titlu": titlu, "url": url,
        "processed_at": datetime.datetime.now().isoformat(timespec="seconds"),
    }
    json.dump(manifest, open(MANIFEST, "w", encoding="utf-8"), ensure_ascii=False, indent=2)

    print(f"Adăugat rând {new_row}: {titlu}")
    print(f"Slug: {slugify(titlu)}")

    subprocess.run(["py", f"{ROOT}/scripts/recalc_rezumat.py"], check=True)

if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Utilizare: py add_youtube_article.py <fisier.txt> <titlu articol> [sursa]")
        sys.exit(1)
    main(sys.argv[1], sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else "YouTube")
