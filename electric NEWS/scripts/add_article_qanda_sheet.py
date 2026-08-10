import openpyxl, os
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = f"{ROOT}/electric-news-calendar-editorial.xlsx"

# (Slug articol, [(Întrebare, Răspuns) x 5])
# Model — câte un articol pilot din fiecare categorie. Întrebările sunt specifice
# conținutului real al articolului, nu generice de categorie (acelea sunt în foaia "Q&A").
ARTICLE_QA = [
    ("ghid-ce-presupune-un-tablou-electric-modern-conform-normativului-i7", [
        ("Ce înseamnă „selectivitate” la protecția diferențială?",
         "Structura pe două niveluri (diferențial general + diferențiale de 30mA pe grupe) în care o scurgere pe un singur circuit oprește doar acel grup, nu tot tabloul."),
        ("De ce e greșit să pui zeci de prize pe o singură siguranță de 16A?",
         "Pentru că un singur scurtcircuit pe oricare aparat conectat scoate din funcțiune tot etajul, nu doar circuitul afectat."),
        ("Am nevoie de protecție la supratensiuni (SPD) dacă am panouri fotovoltaice?",
         "Da — articolul o recomandă explicit pentru case cu echipamente sensibile sau sistem fotovoltaic conectat la tablou."),
        ("De ce contează spațiul liber în tablou, dacă acum nu am nevoie de el?",
         "Ca o extindere ulterioară (aer condiționat nou, stație de încărcare auto) să nu însemne înlocuirea întregului tablou de la zero."),
        ("Cum verific rapid dacă tabloul meu chiar respectă I7?",
         "Testul din articol: dacă electricianul nu poate explica de ce fiecare circuit are exact acea siguranță, tabloul nu e conform, indiferent cum arată."),
    ]),
    ("zigbee-vs-matter-vs-knx-ce-protocol-alegi-pentru-casa-ta-smart-in-2026", [
        ("Care protocol are cele mai multe produse disponibile pe piață?",
         "Zigbee, cu cea mai mare acoperire de produse (prize, becuri, senzori) de la mai mulți producători."),
        ("De ce a apărut Matter dacă Zigbee există deja?",
         "Ca să rezolve fragmentarea Zigbee — un standard unificat, susținut de Google, Apple, Amazon și Samsung, compatibil direct cu asistenții vocali."),
        ("KNX merită pentru o casă deja locuită?",
         "Nu de obicei — KNX cere cablare dedicată din faza de construcție sau renovare majoră, spre deosebire de Zigbee/Matter, care sunt wireless."),
        ("Ce ar trebui să cumpăr dacă vreau compatibilitate pe termen lung?",
         "Dispozitive cu certificare Matter — recomandarea explicită a articolului pentru orice achiziție nouă."),
        ("KNX are vreun avantaj real față de soluțiile wireless?",
         "Fiabilitate maximă fără interferențe radio și control centralizat pentru automatizări complexe, util mai ales combinat cu monitorizare fotovoltaică sau de stocare."),
    ]),
    ("cat-costa-un-sistem-fotovoltaic-complet-in-romania-in-2026-on-grid-hibrid-cu-far", [
        ("Care e diferența de cost între on-grid și hibrid?",
         "Hibridul costă vizibil mai mult, în primul rând din cauza invertorului hibrid, mai scump decât unul standard, chiar înainte de a adăuga bateria."),
        ("Cât adaugă bateria la prețul total al sistemului?",
         "Orientativ 2.000-4.000 euro pentru o baterie de 5-10 kWh, în funcție de brand și capacitate."),
        ("Ce factor influențează cel mai mult prețul final, dincolo de tipul sistemului?",
         "Complexitatea acoperișului — înclinație, orientare, mai multe planuri de montaj cresc semnificativ manopera."),
        ("În cât timp se amortizează un sistem on-grid simplu?",
         "Orientativ 3-8 ani, restul perioadei de funcționare (peste 25 de ani garantați) fiind economie netă."),
        ("De ce nu pot afla un preț exact fără vizită tehnică?",
         "Cifrele din articol sunt repere de piață; costul real depinde de acoperiș și consum propriu, calculabile corect doar la o vizită tehnică."),
    ]),
    ("cat-costa-o-baterie-de-stocare-in-2026-comparatie-deye-huawei-dah-solar", [
        ("Care e cea mai ieftină opțiune la 5 kWh?",
         "Deye, la aproximativ 3.200 de lei."),
        ("Cât reduce subvenția AFM costul unei baterii de 10 kWh?",
         "O baterie BYD de circa 16.460 de lei ajunge la aproximativ 6.460 de lei după o subvenție AFM de 10.000 de lei."),
        ("De ce Huawei costă de trei ori mai mult decât Deye la aceeași capacitate?",
         "Diferența vine din calitatea celulelor, durata garanției și integrarea cu ecosistemul de invertoare, nu din capacitatea nominală."),
        ("Cât de repede se amortizează o baterie cu și fără subvenție?",
         "3,5 ani fără subvenție AFM, sau doar 1,4 ani cu subvenția activă."),
        ("Ce buget total are programul AFM pentru baterii?",
         "400 de milioane de lei, dedicat prosumatorilor care au deja un sistem fotovoltaic instalat."),
    ]),
    ("tariful-de-transport-energie-binom-explicat-pe-intelesul-tuturor", [
        ("Ce se schimbă concret în factură prin BINOM?",
         "O parte din cost depinde de puterea maximă (kW) pe care o poți trage din rețea, nu doar de energia efectiv consumată."),
        ("Cine spune că BINOM nu e o taxă suplimentară?",
         "ACUE, federația operatorilor de distribuție — susține că venitul reglementat total rămâne neschimbat, doar redistribuit altfel."),
        ("De ce sunt prosumatorii dezavantajați de BINOM, potrivit APCE?",
         "Au nevoie de racordare la capacitate mare chiar dacă consumul net e redus datorită producției proprii, deci plătesc pentru capacitate, nu pentru consum real."),
        ("Modelul BINOM funcționează la fel de bine pe orice rețea?",
         "Nu — funcționează bine pe rețele deja modernizate; aplicat pe infrastructură veche din anii 1960-70, riscă să taxeze consumatorii pentru limitările rețelei."),
        ("Ce ar trebui să verific concret pe propria factură?",
         "Structura noii facturi, nu doar totalul — componenta de capacitate poate cântări mai mult decât pare la prima vedere."),
    ]),
    ("unde-montezi-bateria-si-invertorul-fotovoltaic-camerele-interzise-si-capcana-tem", [
        ("De ce e interzisă baia pentru montarea invertorului?",
         "Aburul și stropii corodează contactele electrice, chiar dacă echipamentul are un grad de protecție IP ridicat."),
        ("Care e intervalul ideal de temperatură pentru o baterie cu litiu?",
         "20-25°C — sub 5-7°C și peste 40-45°C, degradarea celulelor accelerează vizibil."),
        ("O baterie de 5 kW livrează mereu 5 kW?",
         "Nu — iarna, la 0-5°C, livrează realist 4-4,2 kW, din cauza rezistenței interne crescute la frig."),
        ("Ce se întâmplă cu bateria vara, la soare direct?",
         "Poate ajunge la 50-60°C, iar BMS-ul reduce singur curentul ca s-o protejeze, exact când ai nevoie de putere maximă."),
        ("Poate garanția bateriei să fie refuzată din cauza montajului greșit?",
         "Da — unele garanții pot fi refuzate pe baza logurilor de temperatură salvate de BMS, la expuneri repetate la temperaturi extreme."),
    ]),
    ("solar-electric-panel-sep-cine-suntem-si-cum-poti-cere-o-oferta", [
        ("Ce fel de proiecte face SEP, dincolo de panourile fotovoltaice?",
         "Baterii de stocare și instalații electrice complete, de la tablouri conforme normativului I7 până la sisteme hibride."),
        ("Evaluarea inițială costă ceva?",
         "Nu — fiecare proiect pornește cu o evaluare gratuită, fără obligații."),
        ("Conținutul editorial de pe Electric NEWS vine din aceeași echipă care face lucrările?",
         "Da — articolul spune explicit că nu publică nimic ce nu ar verifica și în lucrările efective pentru clienți."),
        ("SEP lucrează doar cu persoane fizice?",
         "Nu — lucrează și cu proiecte mai mari, unde e nevoie de o echipă coordonată pe mai multe zone ale țării."),
        ("Ce primesc concret dacă completez formularul de ofertă?",
         "O evaluare reală a situației concrete, nu un preț estimativ vag sau un răspuns automat."),
    ]),
]

def main():
    wb = openpyxl.load_workbook(XLSX)
    if "Q&A Articole" in wb.sheetnames:
        del wb["Q&A Articole"]
    ws = wb.create_sheet("Q&A Articole")

    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
    header_fill = PatternFill(fill_type="solid", fgColor="FF12213A")
    body_font = Font(name="Arial", size=10)
    align_wrap = Alignment(wrap_text=True, vertical="top")
    thin = Side(style="thin", color="FFE5E9EE")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers = ["Slug articol", "Întrebare", "Răspuns"]
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.font, c.fill, c.alignment, c.border = header_font, header_fill, align_wrap, border

    row = 2
    total = 0
    for slug, qa_list in ARTICLE_QA:
        for intrebare, raspuns in qa_list:
            ws.cell(row=row, column=1, value=slug)
            ws.cell(row=row, column=2, value=intrebare)
            ws.cell(row=row, column=3, value=raspuns)
            for col in range(1, 4):
                cell = ws.cell(row=row, column=col)
                cell.font = body_font
                cell.alignment = align_wrap
                cell.border = border
            row += 1
            total += 1

    ws.column_dimensions["A"].width = 55
    ws.column_dimensions["B"].width = 46
    ws.column_dimensions["C"].width = 60
    ws.freeze_panes = "A2"

    wb.save(XLSX)
    print(f"Foaie Q&A Articole creată/actualizată: {total} întrebări, {len(ARTICLE_QA)} articole.")

if __name__ == "__main__":
    main()
