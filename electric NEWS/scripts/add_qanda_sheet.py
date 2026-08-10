import openpyxl, os
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = f"{ROOT}/electric-news-calendar-editorial.xlsx"

# (Categorie, Întrebare, Răspuns, "slug1,slug2,..." sau "pagina:slug")
QA = [
    # ---- Electric & instalații ----
    ("Electric & instalații", "De ce trebuie să-mi schimb tabloul electric dacă încă „merge”?",
     "Un tablou vechi, cu siguranțe fuzibile sau fără protecție diferențială (RCD), nu mai face față consumului modern și nu te protejează la electrocutare — „merge” nu înseamnă „e sigur”.",
     "ghid-ce-presupune-un-tablou-electric-modern-conform-normativului-i7,siguranta-electrica-cele-mai-frecvente-cauze-de-incendii-din-instalatii-vechi"),
    ("Electric & instalații", "Cât costă schimbarea instalației electrice la o casă?",
     "Depinde de suprafață, cât trebuie spart din pereți și calitatea materialelor — manopera e, de regulă, cea mai mare parte din cost, nu materialele.",
     "cat-costa-inlocuirea-instalatiei-electrice-la-o-casa-ghid-de-manopera-pe-etape"),
    ("Electric & instalații", "Am nevoie de autorizație ca să-mi schimb tabloul electric?",
     "Lucrarea trebuie făcută de un electrician autorizat ANRE, dar nu ai nevoie de autorizație de construcție pentru simpla înlocuire a tabloului.",
     "tabloul-electric-autorizatie-inaltime-de-montaj-si-cand-e-semn-de-pericol-daca-s"),
    ("Electric & instalații", "La ce înălțime se montează corect un tablou electric?",
     "Standard, între 1,5 și 1,8 metri de la podea, la loc accesibil, nu ascuns în spatele mobilei.",
     "tabloul-electric-autorizatie-inaltime-de-montaj-si-cand-e-semn-de-pericol-daca-s"),
    ("Electric & instalații", "Ce e RCD/DDR și de ce apare mereu în discuție?",
     "E disjunctorul diferențial care oprește curentul instant dacă „scapă” energie pe un traseu greșit (de exemplu prin corpul tău) — protecția de bază împotriva electrocutării.",
     "rcd-ddr-explicat-diferenta-dintre-protectia-la-atingere-directa-si-indirecta"),
    ("Electric & instalații", "De ce sare siguranța la tabloul electric, tot timpul, la aceleași aparate?",
     "De obicei circuitul e supraîncărcat (prea mulți consumatori pe aceeași linie) sau un aparat are un defect de izolație — nu se rezolvă „băgând o siguranță mai mare”.",
     "tabloul-electric-autorizatie-inaltime-de-montaj-si-cand-e-semn-de-pericol-daca-s"),
    ("Electric & instalații", "E periculoasă o instalație electrică veche, din aluminiu?",
     "Da — casele construite între 1960-1990 au frecvent cablaj din aluminiu, care se oxidează la conexiuni și crește riscul de supraîncălzire și incendiu.",
     "siguranta-electrica-cele-mai-frecvente-cauze-de-incendii-din-instalatii-vechi"),
    ("Electric & instalații", "Ce pot face eu, ca amator, la instalația electrică din casă, fără electrician?",
     "Lucruri minore, vizibile și reversibile (schimbat un bec, o priză simplă, fără modificare de circuit) — orice atinge tabloul, cablajul îngropat sau protecțiile e strict pentru electrician autorizat.",
     "ce-poti-face-diy-la-instalatia-electrica-si-ce-este-strict-interzis-fara-autoriz"),
    ("Electric & instalații", "Câți ani rezistă o instalație electrică înainte să trebuiască schimbată?",
     "Nu există un termen fix legal, dar instalațiile de peste 25 de ani, mai ales fără protecție diferențială, sunt recomandate spre înlocuire de majoritatea electricienilor.",
     "siguranta-electrica-cele-mai-frecvente-cauze-de-incendii-din-instalatii-vechi,cat-costa-inlocuirea-instalatiei-electrice-la-o-casa-ghid-de-manopera-pe-etape"),
    ("Electric & instalații", "Cum îmi dau seama dacă am nevoie urgentă de un electrician, nu doar „la un moment dat”?",
     "Semnale de alarmă: miros de ars fără sursă vizibilă, prize calde la atingere, pâlpâit de lumini fără motiv, siguranțe care sar des — orice din astea nu se amână.",
     "tabloul-electric-autorizatie-inaltime-de-montaj-si-cand-e-semn-de-pericol-daca-s,siguranta-electrica-cele-mai-frecvente-cauze-de-incendii-din-instalatii-vechi"),

    # ---- Smart Home & Automatizări ----
    ("Smart Home & Automatizări", "Ce înseamnă, de fapt, „casă smart”? Nu e doar un bec controlat din telefon?",
     "E despre conectarea mai multor sisteme ale casei (iluminat, climă, securitate, prize) ca să comunice între ele și să poată fi controlate automat sau de la distanță, nu doar un singur gadget izolat.",
     "casa-smart-de-la-zero-ce-inseamna-cat-costa-si-cum-incepi"),
    ("Smart Home & Automatizări", "Cât costă să-mi fac o casă smart de la zero?",
     "De la câteva sute de lei pentru un kit de bază (2-3 prize/becuri inteligente), până la câteva mii de euro pentru automatizare completă cu instalator — depinde total de cât de mult automatizezi.",
     "casa-smart-de-la-zero-ce-inseamna-cat-costa-si-cum-incepi"),
    ("Smart Home & Automatizări", "Trebuie să sparg pereți sau să refac instalația ca să am o casă smart?",
     "Nu neapărat — multe soluții moderne funcționează wireless (Zigbee, Wi-Fi, Matter), montate peste instalația existentă, fără cablare suplimentară.",
     "casa-smart-de-la-zero-ce-inseamna-cat-costa-si-cum-incepi,zigbee-vs-matter-vs-knx-ce-protocol-alegi-pentru-casa-ta-smart-in-2026"),
    ("Smart Home & Automatizări", "Ce diferență e între Zigbee, Matter și Wi-Fi pentru dispozitivele smart?",
     "Zigbee și Matter sunt protocoale dedicate, mai eficiente și mai sigure pentru multe dispozitive mici; Wi-Fi-ul e mai simplu de instalat, dar consumă mai multă rețea și baterie pe termen lung.",
     "zigbee-vs-matter-vs-knx-ce-protocol-alegi-pentru-casa-ta-smart-in-2026"),
    ("Smart Home & Automatizări", "Pot să-mi fac treptat casa smart, sau trebuie totul deodată?",
     "Poți instala pe etape — de exemplu întâi iluminatul, apoi climatizarea, apoi securitatea — majoritatea sistemelor moderne sunt gândite tocmai pentru extindere ulterioară.",
     "casa-smart-de-la-zero-ce-inseamna-cat-costa-si-cum-incepi"),
    ("Smart Home & Automatizări", "E complicat de folosit un sistem smart pentru cineva care nu e „tehnic”?",
     "Configurarea inițială cere puțină răbdare, dar folosirea zilnică (aplicație, comenzi vocale) e gândită să fie la fel de simplă ca orice altă aplicație de telefon.",
     "casa-smart-de-la-zero-ce-inseamna-cat-costa-si-cum-incepi"),
    ("Smart Home & Automatizări", "Casa smart chiar economisește bani, sau e doar confort?",
     "Ambele — termostate și programări automate de climatizare/iluminat pot reduce vizibil consumul, pe lângă confortul de control de la distanță.",
     "casa-smart-de-la-zero-ce-inseamna-cat-costa-si-cum-incepi,automatizari-utile-pentru-case-cu-panouri-fotovoltaice-cum-optimizezi-consumul"),
    ("Smart Home & Automatizări", "Se poate integra o casă smart cu Alexa, Google Home sau Siri?",
     "Da, majoritatea sistemelor moderne, mai ales cele certificate Matter, se conectează direct cu asistenții vocali populari.",
     "casa-smart-de-la-zero-ce-inseamna-cat-costa-si-cum-incepi,matter-in-2026-ce-produse-noi-compatibile-au-aparut-si-de-ce-conteaza-pentru-cas"),
    ("Smart Home & Automatizări", "O casă smart crește valoarea unei locuințe la vânzare?",
     "Poate fi un plus de atractivitate, mai ales pentru cumpărători tineri, dar nu e un factor major de preț ca izolația sau anul construcției — mai degrabă un „bonus” vizibil la vizionare.",
     "casa-smart-de-la-zero-ce-inseamna-cat-costa-si-cum-incepi"),
    ("Smart Home & Automatizări", "Ce se automatizează, concret, într-o casă smart — dincolo de bec și priză?",
     "Iluminat, climatizare, sisteme de umbrire (jaluzele/rulouri), securitate (camere, senzori), echipamente audio-video și contorizare inteligentă a consumului.",
     "casa-smart-de-la-zero-ce-inseamna-cat-costa-si-cum-incepi,automatizari-utile-pentru-case-cu-panouri-fotovoltaice-cum-optimizezi-consumul,camere-de-supraveghere-fara-wifi-si-fara-curent-cu-panou-solar-are-sens-pentru-o"),

    # ---- Fotovoltaic ----
    ("Fotovoltaic", "Cum funcționează, de fapt, panourile fotovoltaice?",
     "Celulele din siliciu transformă lumina soarelui direct în curent continuu, iar un invertor îl transformă apoi în curent alternativ, compatibil cu prizele din casă.",
     "cum-functioneaza-panourile-fotovoltaice-ghid-pentru-incepatori"),
    ("Fotovoltaic", "Cât costă un sistem fotovoltaic complet pentru o casă?",
     "Pentru un sistem rezidențial obișnuit (5-6 kW), prețul complet — echipamente și montaj — se situează în general între 4.000 și 6.000 euro, în funcție de complexitatea acoperișului și echipamentele alese.",
     "cat-costa-un-sistem-fotovoltaic-complet-in-romania-in-2026-on-grid-hibrid-cu-far"),
    ("Fotovoltaic", "Chiar merită investiția, sau se amortizează prea greu?",
     "În condițiile actuale din România, un sistem rezidențial bine dimensionat se amortizează de regulă în 3-8 ani, restul perioadei de funcționare (25+ ani) fiind economie netă.",
     "cat-costa-un-sistem-fotovoltaic-complet-in-romania-in-2026-on-grid-hibrid-cu-far"),
    ("Fotovoltaic", "Panourile produc curent și când e înnorat sau iarna?",
     "Da, dar la capacitate mult redusă — aproximativ 20-30% din producția maximă pe cer înnorat.",
     "cum-functioneaza-panourile-fotovoltaice-ghid-pentru-incepatori"),
    ("Fotovoltaic", "Cât spațiu pe acoperiș am nevoie pentru un sistem rezidențial?",
     "Depinde de puterea dorită, dar orientativ, un sistem de 5-6 kW ocupă aproximativ 25-35 metri pătrați de acoperiș bine orientat.",
     "cum-functioneaza-panourile-fotovoltaice-ghid-pentru-incepatori"),
    ("Fotovoltaic", "Am nevoie de autorizație de construcție ca să montez panouri?",
     "Nu, în general nu e nevoie de autorizație de construcție, atât timp cât montajul nu afectează structura de rezistență a clădirii — dar racordarea la rețea cere documentație separată, prin distribuitor.",
     "cum-functioneaza-panourile-fotovoltaice-ghid-pentru-incepatori,anre-schimba-regulile-de-racordare-garantia-creste-de-la-5-la-20-ce-presupune-pe"),
    ("Fotovoltaic", "Ce se întâmplă cu energia pe care o produc și n-o consum imediat?",
     "Prin compensarea prosumatorilor, surplusul trimis în rețea „scade” din energia cumpărată ulterior, într-o fereastră de timp stabilită prin lege — nu se pierde, dar nici nu se plătește separat ca vânzare.",
     "cum-functioneaza-panourile-fotovoltaice-ghid-pentru-incepatori,compensare-cantitativa-vs-compensare-lunara-ce-inseamna-noua-lege-pentru-factura"),
    ("Fotovoltaic", "Ce întreținere cer panourile fotovoltaice?",
     "Minimă — curățare ocazională cu apă și detergent blând, plus o verificare tehnică anuală a componentelor (invertor, conexiuni) de către un specialist.",
     "cum-functioneaza-panourile-fotovoltaice-ghid-pentru-incepatori"),
    ("Fotovoltaic", "Cât durează, în practică, de la decizie la panouri funcționale pe acoperiș?",
     "De regulă între 2 și 3 luni, incluzând proiectarea, aprobările necesare la distribuitor și montajul propriu-zis.",
     "cum-functioneaza-panourile-fotovoltaice-ghid-pentru-incepatori"),
    ("Fotovoltaic", "Ce durată de viață au panourile fotovoltaice, practic?",
     "Producătorii garantează de obicei 25-30 de ani de funcționare, cu o scădere treptată, dar lentă, a eficienței în acest interval.",
     "cum-functioneaza-panourile-fotovoltaice-ghid-pentru-incepatori"),

    # ---- Baterii & Stocare ----
    ("Baterii & Stocare", "La ce îmi folosește o baterie dacă am deja panouri fotovoltaice?",
     "Stochează surplusul de energie produs ziua, ca să-l poți folosi seara sau noaptea, când panourile nu mai produc — reduce dependența de rețea și de programul de compensare.",
     "sisteme-hibride-on-off-grid-de-ce-creste-cererea-pentru-stocare-in-2026"),
    ("Baterii & Stocare", "Cât costă o baterie de stocare pentru casă?",
     "Variază mult după tehnologie: bateriile LiFePO4, cele mai comune azi pentru case, costă în general între 2.000 și 3.500 lei pe kWh de capacitate.",
     "cat-costa-o-baterie-de-stocare-in-2026-comparatie-deye-huawei-dah-solar"),
    ("Baterii & Stocare", "Ce capacitate de baterie am nevoie pentru o casă obișnuită?",
     "Depinde de consumul serii/nopții, dar orientativ, între 5 și 10 kWh acoperă consumul de bază al unei case obișnuite pentru o seară-noapte.",
     "cat-costa-o-baterie-de-stocare-in-2026-comparatie-deye-huawei-dah-solar"),
    ("Baterii & Stocare", "Cât durează o baterie de stocare până trebuie înlocuită?",
     "Bateriile moderne LiFePO4 au de regulă 10-15 ani de viață utilă și mii de cicluri de încărcare-descărcare, păstrând peste 80% din capacitate chiar și după un deceniu.",
     "bateriile-cu-litiu-explicate-simplu-ghid-pentru-prosumatori,lifepo4-vs-alte-tehnologii-de-baterii-ghid-pentru-cumparatori"),
    ("Baterii & Stocare", "Ce diferență e între o baterie LiFePO4 și una obișnuită (plumb-acid)?",
     "LiFePO4 e mai scumpă la achiziție, dar durează de câteva ori mai mult, e mai sigură termic și suportă mult mai multe cicluri de încărcare complete.",
     "lifepo4-vs-alte-tehnologii-de-baterii-ghid-pentru-cumparatori"),
    ("Baterii & Stocare", "Bateriile de stocare sunt periculoase (explozie, incendiu)?",
     "Bateriile LiFePO4, standardul actual pentru case, sunt considerate printre cele mai sigure tehnologii chimice de baterii, cu risc foarte scăzut de incendiu comparativ cu alte tipuri de litiu.",
     "lifepo4-vs-alte-tehnologii-de-baterii-ghid-pentru-cumparatori,bateriile-cu-litiu-explicate-simplu-ghid-pentru-prosumatori"),
    ("Baterii & Stocare", "Pot adăuga o baterie la un sistem fotovoltaic deja instalat, fără baterie?",
     "În majoritatea cazurilor, da — dacă invertorul existent e compatibil sau poate fi înlocuit cu unul hibrid, adăugarea bateriei ulterior e o extensie obișnuită, nu o reinstalare completă.",
     "sisteme-hibride-on-off-grid-de-ce-creste-cererea-pentru-stocare-in-2026"),
    ("Baterii & Stocare", "Merită o baterie de stocare, sau doar panourile sunt suficiente?",
     "Depinde de programul de consum al casei — dacă folosești curentul mai ales seara, o baterie crește semnificativ gradul de autonomie față de rețea; dacă consumul e mai ales ziua, beneficiul e mai mic.",
     "sisteme-hibride-on-off-grid-de-ce-creste-cererea-pentru-stocare-in-2026"),
    ("Baterii & Stocare", "Ce e o baterie „hibridă” și cum diferă de una „obișnuită”?",
     "Termenul se referă de regulă la invertorul hibrid, capabil să gestioneze simultan panourile, bateria și rețeaua — nu la baterie ca atare, ci la sistemul complet care le leagă inteligent.",
     "sisteme-hibride-on-off-grid-de-ce-creste-cererea-pentru-stocare-in-2026"),
    ("Baterii & Stocare", "Există finanțare de la stat pentru baterii de stocare, separat de panouri?",
     "Da, prin programe dedicate (tip Casa Verde Baterii), care s-au schimbat de la an la an în ce privește bugetul și condițiile — verifică mereu programul activ în anul curent înainte de a aplica.",
     "casa-verde-baterii-2026-buget-de-400-milioane-lei-aprobat-de-guvern-cine-poate-a,programul-afm-baterii-2026-mai-putini-beneficiari-si-posibile-probleme-de-legali"),

    # ---- Reglementări & Prețuri energie ----
    ("Reglementări & Prețuri energie", "Ce înseamnă, mai exact, „prosumator”?",
     "O persoană sau o firmă care consumă curent electric, dar produce și ea, de regulă din panouri fotovoltaice — jumătate consumator, jumătate mic producător.",
     "legea-1602026-a-intrat-in-vigoare-compensare-lunara-pentru-prosumatori-ce-se-sch,cum-citesti-factura-la-energie-ca-prosumator-compensare-tarife-tva"),
    ("Reglementări & Prețuri energie", "Cum funcționează compensarea pentru prosumatori?",
     "Energia produsă și trimisă în rețea se scade din energia consumată, într-o fereastră de timp stabilită prin lege — surplusul rămas necompensat într-o lună se reportează pentru lunile următoare.",
     "compensare-cantitativa-vs-compensare-lunara-ce-inseamna-noua-lege-pentru-factura,legea-1602026-a-intrat-in-vigoare-compensare-lunara-pentru-prosumatori-ce-se-sch"),
    ("Reglementări & Prețuri energie", "De ce plătesc taxe (certificate verzi, tarif de transport) chiar dacă produc și eu curent?",
     "Componentele reglementate din factură (transport, distribuție, certificate verzi) se aplică la energia consumată din rețea, indiferent dacă ești sau nu prosumator — ele finanțează infrastructura comună, nu doar producția.",
     "tariful-de-transport-energie-binom-explicat-pe-intelesul-tuturor,de-ce-platim-certificate-verzi-in-factura-de-energie-si-cat-timp-vor-mai-exista"),
    ("Reglementări & Prețuri energie", "Ce e ANRE și ce rol are în factura mea?",
     "Autoritatea Națională de Reglementare în Energie — stabilește regulile pieței de energie, aprobă tarifele reglementate și verifică respectarea legii de către furnizori și distribuitori.",
     "anre-pregateste-metodologia-noua-pentru-vanzarea-energiei-de-la-prosumatori-pana,anre-reglementeaza-comunitatile-de-energie-si-notificarea-la-80-din-consumul-lun"),
    ("Reglementări & Prețuri energie", "Pot să-mi schimb furnizorul de energie oricând?",
     "Da, schimbarea furnizorului e un drept al consumatorului, gratuit, și se poate face oricând, comparând ofertele disponibile prin platforma oficială a ANRE.",
     "anre-a-retras-prima-licenta-unui-furnizor-din-istoria-institutiei-grenerg-ce-tre"),
    ("Reglementări & Prețuri energie", "Ce se întâmplă dacă furnizorul meu de energie dă faliment sau i se retrage licența?",
     "Ești redistribuit automat către furnizorul de ultimă instanță, fără întreruperea alimentării cu curent — apoi poți alege oricând alt furnizor, dacă vrei.",
     "anre-a-retras-prima-licenta-unui-furnizor-din-istoria-institutiei-grenerg-ce-tre"),
    ("Reglementări & Prețuri energie", "Ce e tariful de transport BINOM și de ce apare separat pe factură?",
     "E componenta din factură care acoperă costul transportului energiei prin rețeaua națională de înaltă tensiune, calculată printr-o formulă distinctă de cea a distribuției locale.",
     "tariful-de-transport-energie-binom-explicat-pe-intelesul-tuturor"),
    ("Reglementări & Prețuri energie", "De ce cresc sau scad tarifele reglementate de la un an la altul?",
     "ANRE recalculează anual componentele reglementate (transport, certificate verzi) pe baza unor formule legale, în funcție de consum estimat, investiții în rețea și evoluția pieței de energie.",
     "certificate-verzi-2027-cum-a-estimat-anre-impactul-in-factura-consumatorilor"),
    ("Reglementări & Prețuri energie", "Am nevoie de contor inteligent ca să fiu prosumator?",
     "În general da — contorul inteligent (bidirecțional) e cel care măsoară separat energia consumată și cea injectată în rețea, esențial pentru calculul corect al compensării.",
     "cum-citesti-factura-la-energie-ca-prosumator-compensare-tarife-tva"),
    ("Reglementări & Prețuri energie", "Unde verific dacă o informație despre legislația energetică e reală, nu doar zvon?",
     "Sursa oficială e mereu site-ul ANRE (anre.ro), unde se publică toate comunicatele, ordinele și metodologiile în vigoare — orice altă sursă ar trebui verificată prin comparație cu acesta.",
     "anre-a-retras-prima-licenta-unui-furnizor-din-istoria-institutiei-grenerg-ce-tre,pagina:sursele-noastre"),

    # ---- YouTube ----
    ("YouTube", "Ce fel de conținut găsesc în categoria YouTube pe Electric NEWS?",
     "Rezumate editoriale scurte (2-3 minute de citit) ale unor videoclipuri relevante despre electric, fotovoltaic, smart home și baterii, de la canale românești de specialitate.",
     "instalatie-electrica-pentru-casa-noua-ce-nu-trebuie-ratat-inainte-de-turnarea-be,unde-montezi-bateria-si-invertorul-fotovoltaic-camerele-interzise-si-capcana-tem"),
    ("YouTube", "De ce citesc un rezumat în loc să mă uit direct la videoclip?",
     "Rezumatul scoate esențialul din 10-20 de minute de video în câteva paragrafe — util când vrei informația rapid, nu tot procesul filmat.",
     ""),
    ("YouTube", "Videoclipurile rezumate sunt verificate tehnic, sau doar preluate?",
     "Rezumatul urmărește fidel conținutul real al videoclipului, fără a adăuga informații neverificate — dacă ceva nu e clar în video, rămâne neclar și în rezumat, nu se completează cu presupuneri.",
     "pagina:sursele-noastre"),
    ("YouTube", "Pot vedea și videoclipul original, nu doar rezumatul?",
     "Da, fiecare articol din categoria YouTube include link direct către videoclipul original și canalul care l-a publicat.",
     "instalatie-electrica-pentru-casa-noua-ce-nu-trebuie-ratat-inainte-de-turnarea-be"),
    ("YouTube", "Ce canale de YouTube sunt urmărite pentru această secțiune?",
     "Canale românești de specialitate în electric, DIY și instalații fotovoltaice — lista se extinde constant, pe măsură ce apar canale noi relevante.",
     "unde-montezi-bateria-si-invertorul-fotovoltaic-camerele-interzise-si-capcana-tem"),
    ("YouTube", "De ce nu sunt toate videoclipurile de pe canalele respective transformate în articole?",
     "Se aleg doar cele cu relevanță clară pentru cititorii Electric NEWS (fotovoltaic, electric, smart home, baterii) — nu orice video de pe canal, ci cele cu conținut tehnic util.",
     ""),
    ("YouTube", "Pot sugera un videoclip pe care să-l rezumați?",
     "Da, prin formularul „Trimite o știre” — orice link YouTube relevant poate fi propus pentru un rezumat.",
     "pagina:trimite-o-stire"),
    ("YouTube", "Thumbnail-ul de la articol e chiar din videoclipul original?",
     "Da, e imaginea oficială de previzualizare a videoclipului de pe YouTube, nu o poză generată sau aleasă separat.",
     ""),
    ("YouTube", "Cât de recent trebuie să fie un videoclip ca să fie rezumat?",
     "Nu contează neapărat vechimea — contează dacă informația e încă valabilă și utilă, mai ales pentru subiecte tehnice care nu se schimbă rapid (de exemplu tutoriale de instalare).",
     ""),
    ("YouTube", "Rezumatele astea sunt sponsorizate de cineva anume?",
     "Nu implicit — dacă un conținut e sponsorizat, apare marcat explicit cu eticheta „Conținut sponsorizat”, vizibil pe pagina articolului.",
     ""),

    # ---- Colaboratori ----
    ("Colaboratori", "Ce înseamnă categoria „Colaboratori” pe Electric NEWS?",
     "E secțiunea care explică cine operează site-ul — Solar Electric Panel (SEP) — și cum poți colabora sau cere o ofertă direct de la echipa din spatele publicației.",
     "solar-electric-panel-sep-cine-suntem-si-cum-poti-cere-o-oferta"),
    ("Colaboratori", "Electric NEWS e un site media independent sau al unei firme?",
     "E operat de Solar Electric Panel (SEP), o firmă din domeniul fotovoltaic și instalații electrice — asta e declarat explicit în footer și în pagina Colaboratori, nu ascuns.",
     "solar-electric-panel-sep-cine-suntem-si-cum-poti-cere-o-oferta"),
    ("Colaboratori", "Dacă sunt electrician sau instalator, pot deveni partener/colaborator?",
     "Momentan nu există un program formal de parteneriat activ pentru alți instalatori — rămâne o opțiune posibilă pe termen lung, dacă volumul de solicitări o justifică.",
     "solar-electric-panel-sep-cine-suntem-si-cum-poti-cere-o-oferta,pagina:contact"),
    ("Colaboratori", "Cum cer o ofertă de la Solar Electric Panel prin acest site?",
     "Prin formularul dedicat „Cere ofertă gratuită”, disponibil în header și la finalul articolelor din categoriile Fotovoltaic, Baterii & Stocare și Electric & instalații.",
     "pagina:cere-oferta"),
    ("Colaboratori", "Solar Electric Panel lucrează doar în București, sau la nivel național?",
     "SEP are acoperire națională, cu echipă și parteneri în mai multe zone ale țării, nu doar în București, unde e sediul.",
     "solar-electric-panel-sep-cine-suntem-si-cum-poti-cere-o-oferta"),
    ("Colaboratori", "De ce nu sunt separate complet articolele „de presă” de partea comercială SEP?",
     "Modelul e transparent asumat: Electric NEWS produce conținut real și verificat, dar scopul de bază e generarea de contact cu cititorii interesați de servicii SEP — asta e declarat explicit, nu ascuns sub formă de site „neutru”.",
     "solar-electric-panel-sep-cine-suntem-si-cum-poti-cere-o-oferta,pagina:sursele-noastre"),
    ("Colaboratori", "Firmele pot apărea ca sponsori pe Electric NEWS?",
     "Da, conținutul sponsorizat e o opțiune posibilă, dar marcat obligatoriu vizibil cu eticheta „Conținut sponsorizat” — nu se amestecă nemarcat cu conținutul editorial obișnuit.",
     ""),
    ("Colaboratori", "Cum contactez direct echipa Electric NEWS / SEP, nu doar pentru ofertă?",
     "Prin pagina de Contact, cu formular dedicat pentru orice solicitare care nu e strict o cerere de ofertă comercială.",
     "pagina:contact"),
    ("Colaboratori", "Pot trimite eu un articol sau o idee de subiect pentru site?",
     "Da, prin formularul „Trimite o știre” — orice sugestie de subiect, link relevant sau idee brută poate fi propusă spre publicare.",
     "pagina:trimite-o-stire"),
    ("Colaboratori", "De unde știu că informațiile tehnice de pe site nu sunt influențate de interesul comercial al SEP?",
     "Pagina „Sursele noastre” explică transparent metodologia — sursele citate explicit, faptele verificate, iar unde există interes comercial direct (ofertă, sponsorizare), e marcat clar, nu ascuns în conținutul „neutru”.",
     "pagina:sursele-noastre"),
]

def main():
    wb = openpyxl.load_workbook(XLSX)
    if "Q&A" in wb.sheetnames:
        del wb["Q&A"]
    ws = wb.create_sheet("Q&A")

    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
    header_fill = PatternFill(fill_type="solid", fgColor="FF12213A")
    body_font = Font(name="Arial", size=10)
    align_wrap = Alignment(wrap_text=True, vertical="top")
    thin = Side(style="thin", color="FFE5E9EE")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers = ["Categorie", "Întrebare", "Răspuns", "Articole (slug-uri, virgulă; pagina:slug pt. pagini statice)"]
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.font, c.fill, c.alignment, c.border = header_font, header_fill, align_wrap, border

    for i, (cat, intrebare, raspuns, articole) in enumerate(QA, start=2):
        ws.cell(row=i, column=1, value=cat)
        ws.cell(row=i, column=2, value=intrebare)
        ws.cell(row=i, column=3, value=raspuns)
        ws.cell(row=i, column=4, value=articole)
        for col in range(1, 5):
            cell = ws.cell(row=i, column=col)
            cell.font = body_font
            cell.alignment = align_wrap
            cell.border = border

    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 46
    ws.column_dimensions["C"].width = 60
    ws.column_dimensions["D"].width = 40
    ws.freeze_panes = "A2"

    wb.save(XLSX)
    print(f"Foaie Q&A creată/actualizată: {len(QA)} întrebări.")

if __name__ == "__main__":
    main()
