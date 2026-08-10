import openpyxl, os
from copy import copy
from collections import Counter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = f"{ROOT}/electric-news-calendar-editorial.xlsx"

# randurile 9-13 din foaia Rezumat sunt fixe, pentru cele 5 categorii originale;
# randul 14 era blank in layout-ul original si e folosit acum pentru YouTube (categorie noua)
CAT_ORDER = ["Reglementări & Prețuri energie", "Baterii & Stocare", "Fotovoltaic",
             "Electric & instalații", "Smart Home & Automatizări"]

def main():
    wb = openpyxl.load_workbook(XLSX)
    cal = wb["Calendar editorial"]
    rez = wb["Rezumat"]

    rows = []
    for r in range(2, cal.max_row + 1):
        titlu = cal.cell(row=r, column=5).value
        if not titlu:
            continue
        rows.append(dict(
            cat=cal.cell(row=r, column=3).value,
            tip=cal.cell(row=r, column=4).value,
            sursa=cal.cell(row=r, column=6).value or "—",
            url=(cal.cell(row=r, column=7).value or "").strip(),
        ))

    total = len(rows)
    with_url = sum(1 for r in rows if r["url"].startswith("http"))
    own = sum(1 for r in rows if "propriu" in r["url"].lower())
    tbd = total - with_url - own

    rez["B3"] = total
    rez["B5"] = with_url
    rez["B6"] = own
    rez["B7"] = tbd

    cat_counts = Counter(r["cat"] for r in rows)
    for i, cat in enumerate(CAT_ORDER):
        rez.cell(row=9 + i, column=2).value = cat_counts.get(cat, 0)
    if cat_counts.get("YouTube", 0) > 0:
        style_ref_A, style_ref_B = rez["A13"], rez["B13"]
        a14, b14 = rez.cell(row=14, column=1), rez.cell(row=14, column=2)
        a14.value, b14.value = "YouTube", cat_counts["YouTube"]
        a14.font, b14.font = copy(style_ref_A.font), copy(style_ref_B.font)
        a14.alignment, b14.alignment = copy(style_ref_A.alignment), copy(style_ref_B.alignment)

    tip_counts = Counter(r["tip"] for r in rows)
    rez["B16"] = tip_counts.get("Agregat", 0)
    rez["B17"] = tip_counts.get("Propriu", 0)

    sursa_counts = Counter(r["sursa"] for r in rows)
    style_ref_A, style_ref_B = rez["A20"], rez["B20"]
    for r in range(20, 20 + max(len(sursa_counts) + 5, 30)):
        rez.cell(row=r, column=1).value = None
        rez.cell(row=r, column=2).value = None
    for i, (sursa, cnt) in enumerate(sorted(sursa_counts.items(), key=lambda x: (-x[1], x[0]))):
        row = 20 + i
        a, b = rez.cell(row=row, column=1), rez.cell(row=row, column=2)
        a.value, b.value = sursa, cnt
        a.font, b.font = copy(style_ref_A.font), copy(style_ref_B.font)
        a.alignment, b.alignment = copy(style_ref_A.alignment), copy(style_ref_B.alignment)

    wb.save(XLSX)
    print(f"Rezumat recalculat: total={total} url={with_url} propriu={own} tbd={tbd} surse={len(sursa_counts)}")

if __name__ == "__main__":
    main()
