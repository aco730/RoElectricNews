import openpyxl, os
from copy import copy
from openpyxl.styles import PatternFill, Font, Alignment

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = f"{ROOT}/electric-news-calendar-editorial.xlsx"
ARTICLE_TEXT_DIR = f"{ROOT}/articole-text"

PUBLICAT = "Publicat"
IN_ASTEPTARE = "În așteptare"

DIACRITICE = {"ă":"a","â":"a","î":"i","ș":"s","ş":"s","ț":"t","ţ":"t"}
def slugify(s):
    out = "".join(DIACRITICE.get(c, c) for c in s.lower())
    keep = [c if c.isalnum() else (" " if c in " -_/" else "") for c in out]
    out = "".join(keep)
    return "-".join(out.split())[:80].strip("-")

def main():
    wb = openpyxl.load_workbook(XLSX)
    ws = wb["Calendar editorial"]

    if ws.cell(row=1, column=8).value != "Status":
        ref_header = ws.cell(row=1, column=7)
        h = ws.cell(row=1, column=8)
        h.value = "Status"
        h.font, h.fill = copy(ref_header.font), copy(ref_header.fill)
        h.border, h.alignment = copy(ref_header.border), copy(ref_header.alignment)
        ws.column_dimensions['H'].width = ws.column_dimensions['G'].width

    fill_publicat = PatternFill(fill_type="solid", fgColor="FF2E7D32")
    font_publicat = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
    fill_asteptare = PatternFill(fill_type="solid", fgColor="FFD6D8DB")
    font_asteptare = Font(name="Arial", size=10, italic=True, color="FF5F6368")
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    seen_slugs = {}
    counts = {PUBLICAT: 0, IN_ASTEPTARE: 0}
    for r in range(2, ws.max_row + 1):
        titlu = ws.cell(row=r, column=5).value
        if not titlu:
            continue
        base_slug = slugify(titlu)
        slug = base_slug
        n = 2
        while slug in seen_slugs:
            slug = f"{base_slug}-{n}"
            n += 1
        seen_slugs[slug] = True

        has_content = os.path.exists(f"{ARTICLE_TEXT_DIR}/{slug}.txt")
        status = PUBLICAT if has_content else IN_ASTEPTARE
        counts[status] += 1

        cell = ws.cell(row=r, column=8)
        cell.value = status
        cell.fill = fill_publicat if has_content else fill_asteptare
        cell.font = font_publicat if has_content else font_asteptare
        cell.alignment = align

    wb.save(XLSX)
    print(f"Status setat: {counts[PUBLICAT]} Publicat, {counts[IN_ASTEPTARE]} În așteptare.")

if __name__ == "__main__":
    main()
